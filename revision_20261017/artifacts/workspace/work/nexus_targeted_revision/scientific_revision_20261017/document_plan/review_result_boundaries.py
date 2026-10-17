from pathlib import Path
import ast
import json
import re

from docx import Document

from main_result_replacements import TARGETS, inspect, local_segment
from update_si_results import PACKET, digest, write_json
from word_edits import accepted


HERE = Path(__file__).resolve().parent
CHANGED_PREFIXES = {
    "abstract": "Climate mitigation scenarios that treat",
    "costs": "The three representative systems",
    "carbon": "The cost difference",
    "capacity": "CCU-GH-FT extends the operating comparison",
    "gtl_operation": "GTL’s average utilization",
}
UNCHANGED_MAIN_TARGETS = {"abstract": 8, "intro": 15, "fig1caption": 23, "spatial": 35, "weeks": 44,
                          "resolution": 68, "solver": 86, "accounting": 65, "figure3caption": 38, "conclusion": 51}


def citation_spans(paragraph):
    spans, offset = [], 0
    for run in paragraph.runs:
        if run.font.strike:
            continue
        end = offset + len(run.text)
        if run.font.superscript and re.search(r"\d", run.text):
            spans.append({"text": run.text, "start": offset, "end": end})
        offset = end
    return spans


def prefixes(path):
    tree = ast.parse(path.read_text(encoding="utf-8"))
    matches = [node.value for node in tree.body if isinstance(node, ast.Assign)
               and any(isinstance(target, ast.Name) and target.id == "prefixes" for target in node.targets)]
    if len(matches) != 1:
        raise ValueError("The response evidence script must contain one literal prefixes mapping")
    result = ast.literal_eval(matches[0])
    for node in tree.body:
        if isinstance(node, ast.For) and isinstance(node.target, ast.Tuple) and [item.id for item in node.target.elts] == ["key", "number", "title"]:
            for key, number, title in ast.literal_eval(node.iter):
                result[key] = ("supp", f"Table S{number} ", f"Table S{number} {title} (excerpt)")
    return result


def main():
    locations = inspect(PACKET)
    citation_records = []
    for name, bound in locations["documents"].items():
        document = Document(PACKET / name)
        for index in TARGETS:
            paragraph = document.paragraphs[bound["paragraph_indices"][str(index)]]
            text = accepted(paragraph)
            segment = local_segment(index, text, bound["language"])
            begin, end = text.index(segment), text.index(segment) + len(segment)
            citations = citation_spans(paragraph)
            intersecting = [citation for citation in citations if citation["start"] < end and citation["end"] > begin]
            if intersecting:
                raise ValueError(f"A planned result segment includes a superscript citation: {name}/P{index}: {intersecting}")
            if citations:
                citation_records.append({"file": name, "clean_reference_index": index, "citations_preserved_outside_edit": citations})
    path = HERE / "refresh_response_evidence.py"
    current = prefixes(path)
    changes = [{"key": key, "before": current[key][1], "after": after} for key, after in CHANGED_PREFIXES.items() if current[key][1] != after]
    for key, index in UNCHANGED_MAIN_TARGETS.items():
        before = locations["documents"]["Manuscript_Clean.docx"]["paragraphs"][str(index)]
        if not before.startswith(current[key][1]):
            raise ValueError(f"The unchanged response prefix no longer matches: {key}")
    documents = {"main": Document(PACKET / "Manuscript_Clean.docx"), "supp": Document(PACKET / "Supplementary_Information.docx")}
    current_only = []
    for key, (role, prefix, description) in current.items():
        count = sum(accepted(paragraph).startswith(prefix) for paragraph in documents[role].paragraphs)
        current_only.append({"key": key, "prefix": prefix, "current_matches": count, "description": description,
                             "will_change_with_result_plan": key in CHANGED_PREFIXES})
    reasons = {8: "Actual cost, service and matching in the Abstract; later-configuration framing respects the carbon-account boundary.",
               15: "Necessary numeric summary; the literature and model-definition portion remains intact.",
               17: "A comparison heading applies to the actual baseline envelope and selected representatives.",
               18: "The actual nondominated set and the three representative configurations are identified separately.",
               19: "New GTL-BH cost and service; actual CCU/GTL cost difference and currency equivalents.",
               20: "New carbon accounts for the named pathways without a lifecycle-reduction claim.",
               21: "New GTL-BH stage contributions and representative-system terminology.",
               23: "Direct supply-chain CI axis and actual service fractions in Figure 1.",
               25: "Saved CCU capacities and actual production.",
               27: "Representative spatial comparisons do not require every configuration to be nondominated.",
               28: "Network descriptors remain consistent with the newly solved configurations.",
               29: "New GTL-BH selected nodes, sites and routes replace the old topology claim.",
               30: "New CCU module/site counts and distances.",
               31: "New DAC-BH flow, capacity and nondominance evidence.",
               33: "Actual service fractions and module terminology in Figure 2.",
               35: "Actual capacity-weighted matching and threshold dependence.",
               36: "The accepted English heading already uses neutral operating terminology.",
               38: "Conditioning-capacity denominator, neutral premium labels and actual totals in Figure 3.",
               40: "Premium dominance is determined from the new schedules.",
               41: "New GTL-BH conditioning utilization, retention and low-load episodes.",
               42: "Preserved GTL values; its comparison with newly solved GTL-BH follows the actual schedules.",
               43: "CCU upstream/synthesis utilization and capacity requirements at the achieved output.",
               44: "Actual six paired fixed-design outcomes and the eighteen-solution verification scope.",
               47: "Only necessary quantitative Discussion facts; the infrastructure-inheritability question is retained.",
               51: "Only the shared-site noun and verb number follow the actual number of co-located airport synthesis sites.",
               65: "Same-configuration combustion and capture-attribution accounts.",
               68: "The route citation is retained; CTL baseline and alternative diagnostic scope are explicit.",
               81: "Scenario-scope description matches both saved CTL diagnostics.",
               86: "Eighteen real certificates, the archived GTL comparator and preserved currency/reference text."}
    report = {"schema_version": 1, "passed": True, "docx_saved": False, "target_paragraphs_per_document": len(TARGETS),
              "citation_spans_crossed_by_edits": 0, "protected_citations": citation_records,
              "evidence_script_sha256": digest(path), "required_prefix_changes": changes,
              "final_result_prefixes": CHANGED_PREFIXES, "paragraph_review": reasons,
              "current_prefix_inventory": current_only,
              "scientific_review": [
                  "The abstract connects near-term and later production configurations without assigning a lifecycle-decarbonization rank to the modelled supply-chain ledger.",
                  "The finite baseline envelope is reconstructed from all thirteen actual cost/intensity points; the three representative configurations are recorded separately.",
                  "Cost and carbon paragraphs use actual production-based accounts and service fractions. No 89-reference subtraction is introduced into plotted baseline intensities.",
                  "Spatial claims use actual selected modules, physical sites and capacity matching; old equal-node-count and zero-H2-route claims are conditional on the new records.",
                  "Operating comparisons use conditioning-capacity denominators for by-product hydrogen and calculate dominance from actual premium components.",
                  "CTL A/B retain the declared diagnostic scope; only the baseline necessary resource bound proves infeasibility with the original 90% requirement.",
                  "Introduction citations, conversion-route citation and exchange-rate citation remain outside the replacement segments. Acknowledgments are outside all plans."
              ]}
    write_json(HERE / "main_result_review.json", report)
    print(json.dumps({"passed": True, "citation_spans_crossed": 0, "required_prefix_changes": changes}, ensure_ascii=False))


if __name__ == "__main__":
    main()
