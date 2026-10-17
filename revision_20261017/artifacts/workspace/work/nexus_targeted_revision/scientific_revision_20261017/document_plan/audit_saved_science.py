from pathlib import Path
import json
import re

from docx import Document

from update_si_results import PACKET, digest, labelled_tables, read_json, rows_of, write_json
from word_edits import accepted


HERE = Path(__file__).resolve().parent
PATTERNS = {
    "old_numeric_values": r"(?<!\d)(?:15\.591\d*|34\.671\d*|23\.83|47\.77|19\.436\d*|63\.24)(?!\d)",
    "old_computation_scope": r"(?:twelve|12)\s+(?:new\s+)?(?:solutions|solves)|six\s+(?:new|recomputed|newly)\s+(?:baseline|configurations|solutions)|seven\s+(?:archived|preserved)|six.{0,45}seven",
    "CTL_infeasibility": r"infeasib|90%|90 percent|90 per cent",
    "carbon_benchmark": r"\b89\b|70%|70 percent|deeper.decarbon|decarbonization group|carbon.diff|avoided.emission",
}


def normalize(text):
    return re.sub(r"\s+", "", text)


def paragraphs_and_cells(document):
    records = [{"location": f"P{index}", "text": accepted(paragraph)}
               for index, paragraph in enumerate(document.paragraphs)]
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                records.append({"location": f"T{table_index}/R{row_index}/C{cell_index}",
                                "text": "\n".join(accepted(p) for p in cell.paragraphs)})
    return records


def audit():
    values = read_json(HERE / "publication_values.json")
    if values["ready"] is not True or values["new_run_count"] != 18:
        raise ValueError("Saved Word audit requires the final eighteen-run publication values")
    for name, expected in values["source_sha256"].items():
        if digest(Path(name)) != expected:
            raise ValueError(f"Publication source changed: {name}")
    result = {"schema_version": 1, "publication_values_sha256": digest(HERE / "publication_values.json"),
              "docx_modified": False, "plans": {}, "tables": {}, "candidates": {}}
    for plan_name in ["main_result_replacements.json", "si_result_replacements.json"]:
        plan = read_json(HERE / plan_name)
        reports = {}
        for name, record in plan["documents"].items():
            document = Document(PACKET / name)
            paragraphs = [normalize(accepted(p)) for p in document.paragraphs]
            missing = [edit["clean_reference_index"] for edit in record["edits"]
                       if normalize(edit["paragraph_after"]) not in paragraphs]
            reports[name] = {"source_sha256": digest(PACKET / name), "expected_replacements": len(record["edits"]),
                             "missing_expected_paragraphs": missing}
        result["plans"][plan_name] = reports
    table_plan = read_json(HERE / "si_result_table_updates.json")
    for name in ["Supplementary_Information.docx", "Supplementary_Information_Marked_Changes.docx"]:
        tables = labelled_tables(Document(PACKET / name))
        result["tables"][name] = {key: rows_of(tables[int(key)]) == rows for key, rows in table_plan["tables"].items()}
    for name in ["Manuscript_Clean.docx", "Supplementary_Information.docx", "Supplementary_Figure_Archive.docx"]:
        records = paragraphs_and_cells(Document(PACKET / name))
        result["candidates"][name] = {key: [row for row in records if re.search(pattern, row["text"], flags=re.I)]
                                          for key, pattern in PATTERNS.items()}
    result["all_prepared_replacements_present"] = all(not report["missing_expected_paragraphs"] for reports in result["plans"].values() for report in reports.values())
    result["all_prepared_tables_match"] = all(all(tables.values()) for tables in result["tables"].values())
    write_json(HERE / "saved_science_audit.json", result)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    audit()
