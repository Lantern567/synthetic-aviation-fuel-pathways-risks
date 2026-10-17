from copy import deepcopy
from pathlib import Path
import argparse
import hashlib
import json
import re

import numpy as np
import pandas as pd
import yaml
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from word_edits import accepted, replace_phrase


HERE = Path(__file__).resolve().parent
WORK = HERE.parent
PACKET = WORK.parents[2] / "Nexus系统提交文件_重算更新20261017"
PAIRED = ["CCU-GH-FT", "CCU-GH-MTJ", "DAC-GH-FT", "DAC-GH-MTJ", "GTL-GH", "CTL"]
BH = ["DAC-BH-FT", "DAC-BH-MTJ", "CCU-BH-FT", "CCU-BH-MTJ", "CTL-BH", "GTL-BH"]
CASES = [(name, label) for name in PAIRED for label in ["baseline", "summer_fixed"]] + [(name, "baseline") for name in BH]
BASELINES = ["CTL", "CTL-BH", "GTL-GH", "GTL", "GTL-BH", "CCU-GH-MTJ", "CCU-GH-FT", "CCU-BH-MTJ", "CCU-BH-FT", "DAC-GH-MTJ", "DAC-GH-FT", "DAC-BH-MTJ", "DAC-BH-FT"]
TABLES = [5, 8, 9, 18, 19, 21, 22, 23, 24, 25, 26, 27, 28, 29, 32]
SOURCE_MAP = {
    5: {"file": "results/{scenario}/baseline/model_parameters.json", "fields": ["Captured-CO2 route-level synthesis electricity"], "additional": "Final accounts apply the common 7.6 kWh/kg FT/MTJ aggregate input"},
    8: {"file": "results/{scenario}/baseline/model_parameters.json", "fields": ["costs.renewable_electricity_cost_yuan_per_mwh"], "additional": "Applied renewable and grid electricity prices"},
    9: {"file": "postprocess/accounts/{scenario}/baseline/accounts.json", "fields": ["process_adjustments.synthesis_energy_kwh_per_kg_fuel"], "additional": "Common aggregate captured-CO2 FT/MTJ synthesis-electricity assumption"},
    18: {"file": "figures_spatial/network_data/network_metrics.csv", "fields": ["节点数量", "h2_supply_module_count", "saf_module_count", "saf_physical_sites", "H2_distance_km", "SAF_distance_km", "关键节点集中度", "关键通道集中度", "跨区域物流比例", "单点失效影响"]},
    19: {"file": "figures_temporal/temporal_data/all_pathway_temporal_metrics.csv", "fields": ["SAFMean", "WorstWeekMean", "ChronicPenaltyPct", "RobustnessPenaltyPct", "TotalPenaltyPct"]},
    21: {"file": "figures_cost/adapted_data/S21_carbon_sensitivity.csv", "fields": ["baseline_g_mj", "with_combustion_g_mj", "with_combustion_without_capture_attribution_g_mj"], "additional": "figures_cost/adapted_data/pathway_summary.csv"},
    22: {"file": "figures_cost/adapted_data/pathway_summary.csv", "fields": ["cost_cny_kg", "fulfillment_ratio"], "additional": "figures_cost/adapted_data/S21_carbon_sensitivity.csv"},
    23: {"file": "figures_temporal/temporal_data/all_pathway_temporal_metrics.csv", "fields": ["SAFMean", "WorstWeekMean", "WorstWeekRetention", "ChronicPenaltyPct", "RobustnessPenaltyPct"]},
    24: {"file": "figures_spatial/network_data/Table_S24_matching.csv", "fields": ["source", "target", "threshold_km", "source_share_percent", "target_share_percent"]},
    25: {"file": "figures_temporal/temporal_data/week_omission_ranges.csv", "fields": ["mean_utilization_pct_min/max", "mean_premium_pct_min/max", "additional_premium_pct_min/max", "retention_pct_min/max"]},
    26: {"file": "figures_resource/updated_joint_week_ranges.json", "fields": ["week", "demand", "solar", "wind", "selected", "inside"]},
    27: {"file": "result_tables/output/solver_certificates.csv", "fields": ["scenario", "label", "rows", "columns", "seconds", "gap"]},
    28: {"file": "results/{scenario}/{label}/summary.json", "fields": ["algebraic_validation"], "additional": "physical_validation.json or byproduct_physical_validation.json"},
    29: {"file": "result_tables/output/collection_report.json", "fields": ["collected_runs", "collected_fixed_asset_pairs", "passed"], "additional": "Independent input, CO2, H2, inventory and physical reports bound to all eighteen summaries"},
    32: {"file": "result_tables/output/new_summary_raw.csv", "fields": ["demand_fulfillment_ratio", "mean_utilization", "worst_week_retention"], "additional": "new_weekly_metrics.csv and fixed_asset_validation.csv"},
}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path, payload):
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


def rows_of(table):
    return [["\n".join(accepted(p) for p in cell.paragraphs) for cell in row.cells] for row in table.rows]


def labelled_tables(document):
    result, caption = {}, None
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = accepted(Paragraph(child, document))
            match = re.match(r"^Table\s+S(\d+)\b", text)
            if match:
                caption = int(match.group(1))
        elif child.tag.endswith("}tbl"):
            if caption in TABLES:
                if caption in result:
                    raise ValueError(f"Duplicate caption binding: Table S{caption}")
                result[caption] = Table(child, document)
            caption = None
    if set(result) != set(TABLES):
        raise ValueError(f"Missing current Word table captions: {set(TABLES) - set(result)}")
    return result


def inspect_sources(packet, snapshot):
    snapshot_data = read_json(snapshot)
    result = {"schema_version": 1, "mode": "table_bindings_only", "snapshot": str(snapshot), "snapshot_sha256": digest(snapshot),
              "snapshot_docx_sha256": snapshot_data["source_sha256"], "required_new_runs": 18, "required_new_baselines": 12,
              "preserved_baseline": "GTL", "documents": {}, "tables": {str(number): SOURCE_MAP[number] for number in TABLES}}
    for filename in ["Supplementary_Information.docx", "Supplementary_Information_Marked_Changes.docx"]:
        path = packet / filename
        document = Document(path)
        tables = labelled_tables(document)
        result["documents"][filename] = {"source_path": str(path), "source_sha256": digest(path),
                                          "tables": {str(number): {"rows": rows_of(table), "columns": len(table.columns),
                                                                             "cell_paragraph_counts": sorted({len(cell.paragraphs) for row in table.rows for cell in row.cells})}
                                                     for number, table in tables.items()}}
    return result


def verify_report(path, summary_hash):
    report = read_json(path)
    hashes = [value for name, value in report.get("input_sha256", {}).items() if name == "summary.json" or name.endswith("/summary.json") or name.endswith("\\summary.json")]
    if report.get("passed") is not True or (report.get("solution_summary_sha256") != summary_hash and hashes != [summary_hash]):
        raise ValueError(f"Independent report does not certify this summary: {path}")
    return report


def byproduct_hydrogen_residual(report):
    keys = [key for key in ["max_h2_node_period_balance_error_kg", "h2_node_period_balance_error_kg"] if key in report["checks"]]
    if len(keys) != 1:
        raise ValueError("The by-product physical report must identify exactly one supported H2 node-period residual field")
    residual = float(report["checks"][keys[0]])
    if not np.isfinite(residual) or residual < 0:
        raise ValueError("The by-product H2 residual must be a finite nonnegative maximum error")
    return residual


def full_gate(results, accounts, result_tables, costs, spatial, temporal):
    collection = read_json(result_tables / "collection_report.json")
    if collection.get("passed") is not True or collection.get("full_revision_set_collected") is not True or collection["collected_runs"] != 18 or collection["collected_baselines"] != 12 or collection["collected_fixed_asset_pairs"] != 6:
        raise ValueError("SI numeric updates require all eighteen accepted runs and six fixed-asset comparisons")
    for name, record in collection["outputs"].items():
        if digest(result_tables / name) != record["sha256"]:
            raise ValueError(f"Result-table contract changed: {name}")
    cost_manifest, network_manifest, time_manifest = read_json(costs / "figure_inputs.json"), read_json(spatial / "provenance.json"), read_json(temporal / "outputs/temporal_verification.json")
    if cost_manifest.get("ready") is not True or cost_manifest["new_run_count"] != 18 or cost_manifest["baseline_pathway_count"] != 13:
        raise ValueError("The complete cost/carbon contract is not ready")
    if network_manifest.get("complete_network_set") is not True or time_manifest.get("full_thirteen_baselines") is not True or time_manifest.get("independent_period_metrics_match") is not True:
        raise ValueError("The complete spatial/temporal contracts are not ready")
    time_input_path = temporal / "inputs/manifest.json"
    if digest(time_input_path) != time_manifest["input_manifest_sha256"]:
        raise ValueError("The verified temporal input manifest changed")
    time_inputs = read_json(time_input_path)
    for name, expected in read_json(costs / "SHA256.json").items():
        if digest(costs / name) != expected:
            raise ValueError(f"Cost contract changed: {name}")
    for name, expected in network_manifest["output_sha256"].items():
        if digest(spatial / name) != expected:
            raise ValueError(f"Spatial contract changed: {name}")
    for name, expected in time_manifest["metric_output_sha256"].items():
        if digest(Path(name)) != expected:
            raise ValueError(f"Temporal metric output changed: {name}")
    runs, physical, co2, tracked = {}, {}, {}, {}
    for folder, manifest in [(costs, read_json(costs / "SHA256.json")), (spatial, network_manifest["output_sha256"])]:
        tracked.update({str(folder / name): expected for name, expected in manifest.items()})
    tracked.update({str(Path(name)): expected for name, expected in time_manifest["metric_output_sha256"].items()})
    tracked.update({str(result_tables / name): record["sha256"] for name, record in collection["outputs"].items()})
    tracked[str(time_input_path)] = digest(time_input_path)
    for record in cost_manifest["sources"]:
        path = Path(record["path"])
        if digest(path) != record["sha256"]:
            raise ValueError(f"Publication source changed: {path}")
        tracked[str(path)] = record["sha256"]
    for record in time_inputs["configurations"].values():
        for name, expected in record["source_hashes"].items():
            if digest(Path(name)) != expected:
                raise ValueError(f"Temporal dispatch source changed: {name}")
            tracked[name] = expected
    for scenario, label in CASES:
        run = results / scenario / label
        path = run / "summary.json"
        if (run / "EXPORT_IN_PROGRESS").exists():
            raise ValueError(f"Run export is incomplete: {run}")
        summary, summary_hash = read_json(path), digest(path)
        if (summary["scenario"], summary["label"]) != (scenario, label) or summary["has_solution"] is not True or not 0 <= summary["gap"] <= .01 or summary["algebraic_validation"]["passed"] is not True:
            raise ValueError(f"Unaccepted source run: {run}")
        account_path = accounts / scenario / label / "accounts.json"
        account = read_json(account_path)
        if account["checks"]["passed"] is not True or account["source_summary_sha256"] != summary_hash:
            raise ValueError(f"Final accounts are not accepted: {account_path}")
        physical_path = run / ("byproduct_physical_validation.json" if "-BH" in scenario else "physical_validation.json")
        physical[scenario, label] = verify_report(physical_path, summary_hash)
        tracked[str(physical_path)] = digest(physical_path)
        if scenario == "GTL-GH":
            ng = physical[scenario, label]["natural_gas_validation"]
            if ng["passed"] is not True or ng["days"] != 28 or ng["periods_per_day"] != 8:
                raise ValueError(f"Incomplete independent natural-gas check: {physical_path}")
        if scenario == "GTL-BH":
            ng_path = run / "extended_natural_gas_validation.json"
            ng = verify_report(ng_path, summary_hash)
            if ng["checks"]["calendar_days"] != 28 or ng["checks"]["periods_per_day"] != 8:
                raise ValueError(f"Incomplete extended natural-gas coverage: {ng_path}")
            tracked[str(ng_path)] = digest(ng_path)
            for name, expected in ng["input_sha256"].items():
                if digest(run / name) != expected:
                    raise ValueError(f"Extended natural-gas input changed: {run / name}")
                tracked[str(run / name)] = expected
        for role in ["hydrogen_transport", "inventory_cost"]:
            report_path = run / f"independent_{role}/{role}_validation.json"
            verify_report(report_path, summary_hash)
            tracked[str(report_path)] = digest(report_path)
        if scenario.startswith("DAC-"):
            co2_path = run / "independent_co2/co2_validation.json"
            co2[scenario, label] = verify_report(co2_path, summary_hash)
            tracked[str(co2_path)] = digest(co2_path)
        expected_weeks = [1, 5, 14, 44] if label == "baseline" else [1, 14, 27, 44]
        if [row["source_week"] for row in summary["weekly"]] != expected_weeks:
            raise ValueError(f"Representative source weeks differ: {run}")
        runs[scenario, label] = summary
        tracked[str(path)] = summary_hash
        tracked[str(account_path)] = digest(account_path)
    for path in [result_tables / "collection_report.json", costs / "figure_inputs.json", spatial / "provenance.json", temporal / "outputs/temporal_verification.json"]:
        tracked[str(path)] = digest(path)
    return runs, physical, co2, tracked


def indexed(path, key, expected):
    frame = pd.read_csv(path)
    if frame[key].duplicated().any() or set(frame[key]) != set(expected):
        raise ValueError(f"Incomplete or duplicate table data: {path}")
    return frame.set_index(key)


def number(value, decimals=2):
    if not np.isfinite(float(value)):
        raise ValueError("A final table cannot contain an unknown numeric value")
    return f"{float(value):.{decimals}f}"


def prepare_rows(binding, results, accounts, result_tables, costs, spatial, temporal, ranges_path):
    runs, physical, co2, tracked = full_gate(results, accounts, result_tables, costs, spatial, temporal)
    network = indexed(spatial / "network_metrics.csv", "Scenario", BASELINES)
    time = indexed(temporal / "temporal_data/all_pathway_temporal_metrics.csv", "Scenario", BASELINES)
    omissions = indexed(temporal / "temporal_data/week_omission_ranges.csv", "scenario", BASELINES)
    costs_table = indexed(costs / "pathway_summary.csv", "scenario", BASELINES)
    carbon = indexed(costs / "S21_carbon_sensitivity.csv", "scenario", BASELINES)
    originals = binding["documents"]["Supplementary_Information.docx"]["tables"]

    def display(name):
        if name == "GTL":
            return "GTL*"
        if name == "CTL" and runs[name, "baseline"]["model_scope"].get("original_minimum_service_constraints_retained") is not True:
            return "CTL†"
        return name

    def order(table):
        return [row[0].rstrip("*†") for row in originals[str(table)]["rows"][1:]]

    tables = {}
    for name in [name for name in BASELINES if name.startswith(("CCU-", "DAC-"))]:
        path = accounts / name / "baseline/accounts.json"
        process = read_json(path)["process_adjustments"]
        if not np.isclose(process["synthesis_energy_kwh_per_kg_fuel"], 7.6, atol=1e-10, rtol=0):
            raise ValueError(f"Captured-CO2 synthesis-energy assumption differs: {name}")
        tracked[str(path)] = digest(path)
    for name in PAIRED:
        path = results / name / "baseline/model_parameters.json"
        if read_json(path)["costs"]["renewable_electricity_cost_yuan_per_mwh"] != 350:
            raise ValueError(f"Applied renewable-electricity price differs: {name}")
        config_path = results / name / "baseline/config.yaml"
        grid_price = yaml.safe_load(config_path.read_text(encoding="utf-8"))["cost_parameters"]["renewable_energy"]["grid_electricity_price_yuan_per_kwh"]
        if not np.isclose(grid_price, .6, atol=1e-10, rtol=0):
            raise ValueError(f"Configured grid-electricity price differs: {name}")
        tracked.update({str(path): digest(path), str(config_path): digest(config_path)})
    for key in [5, 8, 9]:
        tables[key] = deepcopy(originals[str(key)]["rows"])

    def parameter_row(table, label):
        matches = [row for row in tables[table][1:] if row[0] == label]
        if len(matches) != 1:
            raise ValueError(f"Parameter row requires a unique current label: S{table}/{label}")
        return matches[0]

    parameter_row(5, "Direct FT / RWGS-FT")[3] = "Captured-CO₂ FT routes use an aggregate synthesis-electricity input of 7.6 kWh/kg SynAF. The direct GTL/FT comparator retains its configured process-energy coefficient."
    parameter_row(5, "MTJ")[3] = "Captured-CO₂ MTJ routes use the same aggregate synthesis-electricity input of 7.6 kWh/kg SynAF, including the represented conversion stages. Natural-gas MTJ routes retain their separately configured process-energy parameters."
    electricity = parameter_row(8, "Electricity prices")
    electricity[1] = "Renewable electricity 0.35; grid electricity 0.60 CNY/kWh"
    electricity[3] = "Applied operating-cost coefficients for renewable and grid electricity."
    gas_price_path = WORK / "postprocess/natural_gas_price_verification.json"
    gas_prices = read_json(gas_price_path)
    if gas_prices["passed"] is not True or len(gas_prices["checks"]) != 3:
        raise ValueError("Applied natural-gas prices require all three new gas-run ledger checks")
    for row in gas_prices["checks"]:
        if digest(results / row["scenario"] / row["label"] / "summary.json") != row["summary_sha256"] or not np.isclose(row["actual_price_cny_m3"], 2.53, atol=1e-10, rtol=0):
            raise ValueError("Applied natural-gas price evidence differs from its saved result")
    tracked[str(gas_price_path)] = digest(gas_price_path)
    parameter_row(8, "Natural-gas price")[1] = "2.53 CNY/m³"
    parameter_row(9, "Natural-gas feedstock")[1] = "Applied price 2.53 CNY/m³; extraction intensity 0.25; pipeline-transport coefficient 0.01 (configured coefficients)"
    synthesis = ["Captured-CO₂ synthesis electricity", "7.6 kWh/kg SynAF, aggregate route-level input", "CCU and DAC, FT and MTJ configurations", "Common FT/MTJ model assumption used consistently in electricity balances, costs and emissions; no additional 0.8 kWh/kg stage is added."]
    existing = [row for row in tables[9] if row[0] == synthesis[0]]
    if existing:
        if len(existing) != 1:
            raise ValueError("Captured-CO2 electricity parameter row is duplicated")
        existing[0][:] = synthesis
    else:
        tables[9].append(synthesis)
    tables[18] = [["Pathway", "Network\nIDs", "H₂ supply modules", "SynAF modules / sites", "H₂ distance (km)", "SynAF distance (km)", "Top-node share (%)", "Top-route share (%)", "Flow\n>100 km\n(%)", "Largest-source share (%)"]]
    for name in order(18):
        row = network.loc[name]
        tables[18].append([display(name), str(int(row["节点数量"])), str(int(row.h2_supply_module_count)), f"{int(row.saf_module_count)} / {int(row.saf_physical_sites)}",
                           number(row.H2_distance_km) if row.H2_weekly_quantity > 1e-5 else "—", number(row.SAF_distance_km),
                           *[number(row[key]) for key in ["关键节点集中度", "关键通道集中度", "跨区域物流比例", "单点失效影响"]]])
    tables[19] = [["Pathway", "Mean SynAF util. (%)", "Worst-week util. (%)", "Mean-utilization premium (%)", "Additional worst-week premium (%)", "Total premium (%)"]]
    for name in ["GTL-BH", "GTL", "CCU-GH-FT"]:
        row = time.loc[name]
        tables[19].append([display(name), *[number(row[key]) for key in ["SAFMean", "WorstWeekMean", "ChronicPenaltyPct", "RobustnessPenaltyPct", "TotalPenaltyPct"]]])
    tables[21] = [originals["21"]["rows"][0]]
    for name in order(21):
        row, c = costs_table.loc[name], carbon.loc[name]
        if not np.isclose(row.carbon_g_mj, c.baseline_g_mj, rtol=1e-9, atol=1e-8):
            raise ValueError(f"Baseline carbon contracts differ: {name}")
        tables[21].append([display(name), number(row.fulfillment_ratio * 100), number(row.cost_cny_kg), number(c.baseline_g_mj, 3), number(c.with_combustion_g_mj, 3), number(c.with_combustion_without_capture_attribution_g_mj, 3)])
    combustion = carbon.combustion_addition_g_mj.to_numpy()
    if not np.allclose(combustion, combustion[0], atol=1e-9, rtol=0):
        raise ValueError("Uniform combustion sensitivity requires a common configured composition and heating value")
    tables[22] = [originals["22"]["rows"][0]]
    for a, b in [("GTL-BH", "GTL"), ("GTL", "CCU-GH-FT")]:
        x, y = costs_table.loc[a], costs_table.loc[b]
        tables[22].append([f"{display(a)} → {display(b)}", f"{number(x.fulfillment_ratio * 100)}% → {number(y.fulfillment_ratio * 100)}%",
                           f"{number(x.cost_cny_kg, 3)} → {number(y.cost_cny_kg, 3)} CNY/kg of actual output."])
    tables[22].extend([["Service basis", "Achieved baseline output", "Costs and intensities use each configuration’s actual production."],
                       ["Uniform combustion addition", "Same output as each baseline", f"Adds {number(combustion[0], 3)} g CO₂/MJ to each baseline intensity."],
                       ["Capture-attribution sensitivity", "Same output as each baseline", "The increase equals the configuration-specific capture attribution in Table S21."]])
    tables[23] = [["Pathway", "Mean SynAF util. (%)", "Worst-week util. (%)", "Worst-week retention (%)", "Mean-utilization premium (%)", "Additional worst-week premium (%)"]]
    for name in order(23):
        tables[23].append([display(name), *[number(time.loc[name, key]) for key in ["SAFMean", "WorstWeekMean", "WorstWeekRetention", "ChronicPenaltyPct", "RobustnessPenaltyPct"]]])
    matching = pd.read_csv(spatial / "Table_S24_matching.csv").set_index(["source", "target", "threshold_km"])
    tables[24] = [originals["24"]["rows"][0]]
    for a, b in [("GTL-BH", "GTL"), ("GTL-BH", "CCU-GH-FT"), ("GTL", "CCU-GH-FT")]:
        for radius in [1, 10, 25, 50, 100]:
            row = matching.loc[a, b, radius]
            tables[24].append([f"{display(a)} → {display(b)}", str(radius), number(row.source_share_percent), number(row.target_share_percent)])
    tables[25] = [["Pathway", "Mean SynAF util. range (%)", "Mean-utilization premium range (%)", "Additional worst-week premium range (%)", "Retention range (%)"]]
    for name in order(25):
        row = omissions.loc[name]
        tables[25].append([display(name), *[f"{number(row[key + '_min'])}–{number(row[key + '_max'])}" for key in ["mean_utilization_pct", "mean_premium_pct", "additional_premium_pct", "retention_pct"]]])
    ranges = read_json(ranges_path)
    if len(ranges) != 12 or len({row["week"] for row in ranges}) != 12:
        raise ValueError("Joint-range diagnostic must contain the twelve actual candidate weeks")
    tables[26] = [originals["26"]["rows"][0]] + [[str(row["week"]), *[number(row[key], 3) for key in ["demand", "solar", "wind"]], "Yes" if row["selected"] else "No", "Yes" if row["inside"] else "No"] for row in ranges]
    tables[27], tables[28] = [originals["27"]["rows"][0]], [originals["28"]["rows"][0]]
    for name, label in CASES:
        summary, p = runs[name, label], physical[name, label]
        code = f"{display(name)} / {'A' if label == 'baseline' else 'B'}"
        scope_label = "Diagnostic model feasible; global bound certificate" if name == "CTL" else "Full-model feasible; global bound certificate"
        tables[27].append([code, str(summary["rows"]), str(summary["columns"]), number(summary["seconds"]), number(summary["gap"] * 100, 4), scope_label])
        algebraic = summary["algebraic_validation"]
        if "-BH" in name:
            hydrogen = byproduct_hydrogen_residual(p)
            electricity = "—"
        else:
            hydrogen = p["max_hydrogen_balance_residual_kg"]
            electricity = f"{p['max_electricity_deficit_mwh']:.2e}"
        tables[28].append([code, f"{algebraic['max_constraint_absolute_residual']:.2e}", f"{algebraic['max_variable_bound_violation']:.2e}", electricity, f"{hydrogen:.2e}", "Passed"])
    solar, wind = physical["CCU-GH-FT", "baseline"]["input_profile_validation"]["solar"], physical["CCU-GH-FT", "baseline"]["input_profile_validation"]["wind"]
    ctl_scopes = [runs["CTL", label]["model_scope"] for label in ["baseline", "summer_fixed"]]
    if any(scope.get("allow_ctl_shortage") is not True for scope in ctl_scopes):
        raise ValueError("Both CTL results must identify their actual partial-supply diagnostic scope")
    ctl_upper = float(ctl_scopes[0]["baseline_necessary_resource_upper_bound"])
    if not 0 < ctl_upper < .9 or any(scope.get("summer_resource_bound_is_not_an_infeasibility_proof") is not True for scope in ctl_scopes):
        raise ValueError("CTL baseline resource-bound evidence and summer scope must remain separately identified")
    tables[29] = [["Check", "Recorded outcome", "Validation scope"],
                  ["Solar source reconstruction", f"{solar['hdf5_files']:,} HDF5 files; {solar['phases']} phases; {solar['periods']:,} project-period records checked", "Phase-specific DC/AC capacities, 14% system losses, 96% inverter efficiency and UTC+8 alignment."],
                  ["Wind aggregation", f"{wind['selected_phases']} phases; {wind['periods']:,} project-period records checked", "Mean MW and three-hour energy from the cyclic 8,760-hour source sequence."],
                  ["Selected profiles and demand", "A: weeks 1, 5, 14, 44; B: weeks 1, 14, 27, 44; 224 periods each", "Demand and resource inputs selected together and bound to the same saved summaries."],
                  ["Full-model and physical checks", "Eighteen accepted solutions pass", "All rows, bounds and integers in each stated model scope; independent supply, stock, capacity, transport and demand checks."],
                  ["Fixed production assets", "Six A/B pathway pairs pass", "SynAF and H₂ build/capacity; DAC build where present. Inventory and routes are reoptimized."],
                  ["DAC CO₂ accounting", "Six accepted DAC solutions pass", "Node-period, source-week and system-wide balances; capture expenditure, electricity, equipment emissions and transport accounts."],
                  ["H₂ transport accounting", "Eighteen accepted solutions pass", "Four actual weekly contracts, period route capacity, route-specific prices and emissions per 100 km."],
                  ["Inventory accounting", "Eighteen accepted solutions pass", "Holding costs use 672 h in optimization; storage emissions integrate three-hour opening/closing mean stocks over 224 intervals."],
                  ["Natural-gas accounting", "GTL-GH A/B and GTL-BH A pass", "Twenty-eight daily balances, source-technology compatibility, route allocation, truck/LNG capacities and transport/upstream ledgers; airport-integrated procurement is recorded separately."],
                  ["CTL diagnostic scope", "The original 90% minimum-service inequality is removed in A and B", "All other constraints are retained; B fixes the production assets from diagnostic A."],
                  ["CTL baseline resource bound", f"Necessary baseline supply upper bound: {number(100 * ctl_upper, 4)}% < 90%", "The bound establishes infeasibility of baseline A with the original service requirement. It does not establish infeasibility for the summer week set."]]
    tables[32] = [["Pathway", "Week set", "Demand met (%)", "Mean SynAF util. (%)", "Worst-week retention (%)", "Mean-utilization premium (%)", "Additional worst-week premium (%)"]]
    for name in PAIRED:
        for label in ["baseline", "summer_fixed"]:
            row = runs[name, label]
            mean, retained = row["mean_utilization"], row["worst_week_retention"]
            tables[32].append([display(name), "A" if label == "baseline" else "B", number(row["demand_fulfillment_ratio"] * 100), number(mean * 100), number(retained * 100), number((1 / mean - 1) * 100), number((1 / (mean * retained) - 1 / mean) * 100)])
    tracked[str(ranges_path)] = digest(ranges_path)
    return {"schema_version": 1, "ready": True, "required_new_runs": 18, "accepted_new_runs": 18,
            "binding_sha256": hashlib.sha256(json.dumps(binding, ensure_ascii=False, sort_keys=True).encode()).hexdigest(),
            "source_sha256": tracked, "tables": {str(key): value for key, value in tables.items()},
            "caption_dependencies": {"S18_S21_S23": "Only GTL carries the archived-comparator asterisk; CTL carries the partial-supply diagnostic dagger in both week sets.",
                                     "CTL": "A/B remove the original 90% minimum-service inequality and retain the other constraints. B fixes diagnostic A production assets. The necessary resource bound proves only baseline A infeasibility at 90% service.",
                                     "S27_S28": "Eighteen accepted solutions. Seconds denote final full-model dispatch; separate records retain design/projection timings.",
                                     "S28": "Electricity deficits refer to shared renewable budgets; grid-only BH configurations show an em dash."}}


def apply_updates(document, updates, binding, filename, marked=False):
    if updates.get("ready") is not True or updates.get("accepted_new_runs") != 18:
        raise ValueError("Numeric Word edits require a complete accepted result plan")
    binding_hash = hashlib.sha256(json.dumps(binding, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    if updates.get("binding_sha256") != binding_hash:
        raise ValueError("The prepared numeric plan belongs to a different Word table binding")
    for path, expected in updates["source_sha256"].items():
        if digest(Path(path)) != expected:
            raise ValueError(f"Prepared numeric source changed: {path}")
    tables = labelled_tables(document)
    for key, final_rows in updates["tables"].items():
        table = tables[int(key)]
        if rows_of(table) != binding["documents"][filename]["tables"][key]["rows"]:
            raise ValueError(f"Current table was edited after binding: {filename}/S{key}")
        if any(len(row) != len(table.columns) for row in final_rows):
            raise ValueError(f"Table column count differs: S{key}")
        while len(table.rows) < len(final_rows):
            table._tbl.append(deepcopy(table.rows[-1]._tr))
        while len(table.rows) > len(final_rows):
            table._tbl.remove(table.rows[-1]._tr)
        for row, final in zip(table.rows, final_rows):
            for cell, text in zip(row.cells, final):
                if len(cell.paragraphs) != 1:
                    raise ValueError(f"Table cell requires a local paragraph-preserving edit: S{key}")
                paragraph = cell.paragraphs[0]
                before = accepted(paragraph)
                if before != text:
                    replace_phrase(paragraph, before, text, marked=marked)
    return document


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["inspect", "prepare"])
    parser.add_argument("--packet", type=Path, default=PACKET)
    parser.add_argument("--snapshot", type=Path, default=HERE / "current_text/Supplementary_Information.docx.json")
    parser.add_argument("--binding", type=Path, default=HERE / "si_result_table_bindings.json")
    parser.add_argument("--out", type=Path, default=HERE / "si_result_table_updates.json")
    parser.add_argument("--results", type=Path, default=WORK / "results")
    parser.add_argument("--accounts", type=Path, default=WORK / "postprocess/accounts")
    parser.add_argument("--result-tables", type=Path, default=WORK / "result_tables/output")
    parser.add_argument("--costs", type=Path, default=WORK / "figures_cost/adapted_data")
    parser.add_argument("--spatial", type=Path, default=WORK / "figures_spatial/network_data")
    parser.add_argument("--temporal", type=Path, default=WORK / "figures_temporal")
    parser.add_argument("--ranges", type=Path, default=WORK / "figures_resource/updated_joint_week_ranges.json")
    args = parser.parse_args()
    if args.mode == "inspect":
        binding = inspect_sources(args.packet, args.snapshot)
        write_json(args.binding, binding)
        print(json.dumps({"mode": "inspect", "tables": TABLES, "docx_saved": False}))
    else:
        binding = read_json(args.binding)
        plan = prepare_rows(binding, args.results, args.accounts, args.result_tables, args.costs, args.spatial, args.temporal, args.ranges)
        write_json(args.out, plan)
        print(json.dumps({"mode": "prepare", "accepted_new_runs": 18, "tables": TABLES, "docx_saved": False}))


if __name__ == "__main__":
    main()
