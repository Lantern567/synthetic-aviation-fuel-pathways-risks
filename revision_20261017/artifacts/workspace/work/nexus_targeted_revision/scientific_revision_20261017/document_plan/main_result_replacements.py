from pathlib import Path
import argparse
import hashlib
import json
import re

import numpy as np
import pandas as pd
from docx import Document

from word_edits import accepted, replace_phrase
from update_si_results import BASELINES, CASES, PACKET, WORK, digest, full_gate, indexed, read_json, write_json


HERE = Path(__file__).resolve().parent
TARGETS = [8, 15, 17, 18, 19, 20, 21, 23, 25, 27, 28, 29, 30, 31, 33, 35, 36, 38, 40, 41, 42, 43, 44, 47, 51, 65, 68, 81, 86]
DOCUMENTS = {"Manuscript_Clean.docx": "en", "Manuscript_Marked_Changes.docx": "en",
             "作者审阅/Manuscript_Chinese.docx": "zh", "作者审阅/Manuscript_Marked_Changes_Chinese.docx": "zh"}


def inspect(packet):
    documents = {}
    reference = {"en": Document(packet / "Manuscript_Clean.docx"), "zh": Document(packet / "作者审阅/Manuscript_Chinese.docx")}
    for name, language in DOCUMENTS.items():
        path = packet / name
        document = Document(path)
        lookup = {}
        for index, paragraph in enumerate(document.paragraphs):
            text = accepted(paragraph)
            lookup.setdefault(re.sub(r"\s+", "", text), []).append((index, text))
        paragraphs, indices = {}, {}
        for index in TARGETS:
            expected = accepted(reference[language].paragraphs[index])
            normalized = re.sub(r"\s+", "", expected)
            matches = lookup.get(normalized, [])
            if len(matches) != 1:
                raise ValueError(f"Result paragraph is not uniquely matched to the current clean text: {name}/P{index}")
            indices[str(index)], paragraphs[str(index)] = matches[0]
        for index in [15, 51, 68, 81, 86]:
            local_segment(index, paragraphs[str(index)], language)
        documents[name] = {"language": language, "source_path": str(path), "source_sha256": digest(path),
                           "paragraph_indices": indices, "paragraphs": paragraphs}
    return {"schema_version": 1, "mode": "current_result_locations", "docx_saved": False,
            "scope": "Result facts and necessary captions/headings; acknowledgments are outside this plan.",
            "documents": documents}


def fmt(value, digits=2):
    if not np.isfinite(float(value)):
        raise ValueError("Publication facts must be finite")
    return f"{float(value):.{digits}f}"


def integer(value):
    if abs(float(value) - round(float(value))) > 1e-8:
        raise ValueError("A discrete count is not an integer")
    return str(round(float(value)))


def names(items, language="en"):
    if language == "zh":
        return "、".join(items)
    if len(items) < 2:
        return "".join(items)
    if len(items) == 2:
        return " and ".join(items)
    return ", ".join(items[:-1]) + " and " + items[-1]


def facts(results, accounts, result_tables, costs, spatial, temporal):
    runs, physical, co2, tracked = full_gate(results, accounts, result_tables, costs, spatial, temporal)
    cost = indexed(costs / "pathway_summary.csv", "scenario", BASELINES)
    time = indexed(temporal / "temporal_data/all_pathway_temporal_metrics.csv", "Scenario", BASELINES)
    network = indexed(spatial / "network_metrics.csv", "Scenario", BASELINES)
    sensitivity = indexed(costs / "S21_carbon_sensitivity.csv", "scenario", BASELINES)
    matching = pd.read_csv(spatial / "Table_S24_matching.csv")
    if matching.duplicated(["source", "target", "threshold_km"]).any():
        raise ValueError("Capacity matching contains duplicate pairs")
    selected = matching[(matching.source == "GTL") & (matching.target == "CCU-GH-FT")]
    if set(selected.threshold_km) != {1, 10, 25, 50, 100}:
        raise ValueError("GTL to CCU-GH-FT capacity matching is incomplete")
    match25 = selected.set_index("threshold_km").loc[25].to_dict()
    sites = pd.read_csv(spatial / "synthesis_sites.csv")
    shared_sites = sites[sites.scenario.eq("GTL")].merge(sites[sites.scenario.eq("CCU-GH-FT")], on=["latitude", "longitude"], suffixes=("_GTL", "_CCU_GH_FT"))
    if not len(shared_sites):
        raise ValueError("The shared-site interpretation requires an actual co-located synthesis coordinate")
    envelope = []
    for name, row in cost.iterrows():
        dominated = any(other.cost_cny_kg <= row.cost_cny_kg and other.carbon_g_mj <= row.carbon_g_mj
                        and (other.cost_cny_kg < row.cost_cny_kg or other.carbon_g_mj < row.carbon_g_mj)
                        for _, other in cost.iterrows())
        if not dominated:
            envelope.append(name)
    envelope = sorted(envelope, key=lambda name: (cost.loc[name, "carbon_g_mj"], cost.loc[name, "cost_cny_kg"]))
    declared = set(pd.read_csv(costs / "nondominated_envelope.csv").scenario)
    if set(envelope) != declared:
        raise ValueError("Independently reconstructed baseline envelope differs from the figure contract")
    for name in BASELINES:
        if not np.isclose(cost.loc[name, "fulfillment_ratio"], time.loc[name, "DemandFulfillmentRatio"], atol=1e-9, rtol=0):
            raise ValueError(f"Temporal and cost service fractions differ: {name}")
        if name != "GTL" and not np.isclose(time.loc[name, "SAFMean"], runs[name, "baseline"]["mean_utilization"] * 100, atol=1e-8, rtol=0):
            raise ValueError(f"Mean utilization differs from the certified run: {name}")
    carbon_records = {}
    for name in ["GTL-BH", "GTL", "CCU-GH-FT"]:
        path = costs / f"{name}_carbon_emissions.json"
        carbon_records[name] = read_json(path)
        tracked[str(path)] = digest(path)
    ctl_a, ctl_b = runs["CTL", "baseline"], runs["CTL", "summer_fixed"]
    if any(run["model_scope"].get("allow_ctl_shortage") is not True for run in [ctl_a, ctl_b]):
        raise ValueError("CTL A/B must retain the declared partial-supply diagnostic scope")
    upper = float(ctl_a["model_scope"]["baseline_necessary_resource_upper_bound"])
    if not 0 < upper < .9 or ctl_a["model_scope"].get("summer_resource_bound_is_not_an_infeasibility_proof") is not True:
        raise ValueError("CTL baseline infeasibility evidence must not be extended to the summer set")
    bound_path = WORK / "validation/ctl_resource_bounds.json"
    ctl_bounds = read_json(bound_path)
    bounds_by_label = {record["label"]: record for record in ctl_bounds["cases"]}
    if set(bounds_by_label) != {"baseline", "summer_fixed"}:
        raise ValueError("The CTL resource-bound evidence must contain both actual input sequences")
    for label, record in bounds_by_label.items():
        for path, expected in record["input_sha256"].items():
            if digest(Path(path)) != expected:
                raise ValueError(f"CTL resource-bound input changed: {path}")
            tracked[path] = expected
        expected_implication = label == "baseline"
        if record["proves_original_90_percent_requirement_infeasible"] is not expected_implication:
            raise ValueError("CTL baseline/summer infeasibility implications differ from their actual necessary bounds")
    if not np.isclose(bounds_by_label["baseline"]["necessary_resource_supply_fraction_upper_bound"], upper, atol=1e-6, rtol=0):
        raise ValueError("CTL resource bound differs from the accepted summary scope")
    tracked[str(bound_path)] = digest(bound_path)
    return {"runs": runs, "cost": cost, "time": time, "network": network, "sensitivity": sensitivity,
            "frontier": envelope, "match25": match25, "directly_shared_synthesis_sites": json.loads(shared_sites.to_json(orient="records")),
            "matching_constant": bool(np.allclose(selected.matched_capacity_kg_h, match25["matched_capacity_kg_h"], rtol=0, atol=.01)),
            "matching": matching, "temporal_meta": read_json(temporal / "inputs/manifest.json")["configurations"],
            "summary_hashes": {f"{name}/{label}": digest(results / name / label / "summary.json") for name, label in CASES},
            "validation_counts": {"accepted_new_runs": len(runs), "independent_physical_reports": len(physical),
                                  "independent_dac_co2_reports": len(co2), "independent_hydrogen_transport_reports": 18,
                                  "independent_inventory_cost_reports": 18, "fixed_production_asset_pairs": 6,
                                  "independent_natural_gas_schedules": 3},
            "natural_gas_validation": {"GTL-GH/baseline": physical["GTL-GH", "baseline"]["natural_gas_validation"],
                                       "GTL-GH/summer_fixed": physical["GTL-GH", "summer_fixed"]["natural_gas_validation"],
                                       "GTL-BH/baseline": read_json(results / "GTL-BH/baseline/extended_natural_gas_validation.json")},
            "carbon_records": carbon_records, "ctl_upper": bounds_by_label["baseline"]["necessary_resource_supply_fraction_upper_bound"],
            "ctl_bounds": ctl_bounds, "tracked": tracked}


def publication_values(data):
    baselines = {}
    for name in BASELINES:
        row, time, network, meta = data["cost"].loc[name], data["time"].loc[name], data["network"].loc[name], data["temporal_meta"][name]
        baselines[name] = {"cost_cny_kg": float(row.cost_cny_kg), "carbon_g_co2e_mj": float(row.carbon_g_mj),
                           "demand_fulfillment_ratio": float(row.fulfillment_ratio), "production_kg": float(row.production_kg),
                           "synaf_capacity_kg_h": float(meta["synaf_capacity_kg_h"]), "hydrogen_capacity_kg_h": float(meta["hydrogen_capacity_kg_h"]),
                           "hydrogen_utilization_basis": meta["hydrogen_basis"],
                           "temporal_metrics": json.loads(data["time"].loc[[name]].to_json(orient="index", force_ascii=False))[name],
                           "network_metrics": json.loads(data["network"].loc[[name]].to_json(orient="index", force_ascii=False))[name],
                           "carbon_sensitivity": json.loads(data["sensitivity"].loc[[name]].to_json(orient="index", force_ascii=False))[name],
                           "source_kind": meta["source_kind"], "nondominated_baseline_envelope": name in data["frontier"]}
        baselines[name]["temporal_summary"] = {
            "mean_synaf_utilization_fraction": float(time.SAFMean) / 100,
            "mean_hydrogen_utilization_fraction": None if name == "GTL" else float(time.H2Mean) / 100,
            "worst_week_synaf_utilization_fraction": float(time.WorstWeekMean) / 100,
            "worst_week_retention_fraction": float(time.WorstWeekRetention) / 100,
            "mean_utilization_capacity_premium_pct": float(time.ChronicPenaltyPct),
            "additional_worst_week_capacity_premium_pct": float(time.RobustnessPenaltyPct),
            "total_capacity_premium_pct": float(time.TotalPenaltyPct),
            "additional_worst_week_share_pct": float(time.RobustnessSharePct),
            "low_load_episodes_longer_than_12h": int(time.LowLoadEpisodes12h),
            "maximum_low_load_duration_hours": float(time.MaxLowRunHours)}
    pairs = {}
    for name in ["CCU-GH-FT", "CCU-GH-MTJ", "DAC-GH-FT", "DAC-GH-MTJ", "GTL-GH", "CTL"]:
        pairs[name] = {}
        for label in ["baseline", "summer_fixed"]:
            run = data["runs"][name, label]
            mean, retained = run["mean_utilization"], run["worst_week_retention"]
            pairs[name][label] = {key: run[key] for key in ["demand_fulfillment_ratio", "production_kg", "demand_kg", "shortage_kg", "saf_capacity_kg_h", "h2_capacity_kg_h", "mean_utilization", "worst_week_retention", "weekly", "gap", "model_scope"]}
            pairs[name][label].update({"mean_utilization_capacity_premium_pct": (1 / mean - 1) * 100,
                                       "additional_worst_week_capacity_premium_pct": (1 / (mean * retained) - 1 / mean) * 100})
    gaps = [{"scenario": name, "label": label, "gap": data["runs"][name, label]["gap"],
             "gap_percent": data["runs"][name, label]["gap"] * 100,
             "source_summary_sha256": data["summary_hashes"][f"{name}/{label}"]}
            for name, label in CASES]
    return {"schema_version": 1, "ready": True, "baseline_count": 13, "new_run_count": 18,
            "units": {"cost": "CNY/kg actual fuel output", "carbon": "g CO2e/MJ actual fuel output", "capacities": "kg/h",
                      "temporal_summary": "Field suffixes explicitly distinguish fractions, percentages, hours and counts",
                      "temporal_metric_percentage_columns": ["SAFMean", "H2Mean", "WorstWeekMean", "WorstWeekRetention", "ChronicPenaltyPct", "RobustnessPenaltyPct", "TotalPenaltyPct", "RobustnessSharePct"],
                      "temporal_metric_fraction_columns": ["MeanSAFUtilization", "MeanSynAFUtilization", "DemandFulfillmentRatio"],
                      "fixed_asset_comparisons": "mean_utilization and worst_week_retention are fractions; premium fields ending pct are percentages"},
            "baselines": baselines, "actual_nondominated_baselines": data["frontier"],
            "selected_representative_pathways": ["GTL-BH", "GTL", "CCU-GH-FT"],
            "capacity_matching": json.loads(data["matching"].to_json(orient="records", force_ascii=False)),
            "directly_shared_GTL_CCU_GH_FT_synthesis_sites": data["directly_shared_synthesis_sites"],
            "fixed_asset_comparisons": pairs, "gap_certificates": gaps,
            "maximum_total_objective_gap": max(row["gap"] for row in gaps),
            "maximum_total_objective_gap_percent": max(row["gap_percent"] for row in gaps),
            "validation_counts": data["validation_counts"], "ctl_baseline_necessary_resource_upper_bound": data["ctl_upper"],
            "ctl_resource_bounds": data["ctl_bounds"],
            "natural_gas_validation": {name: {key: value for key, value in report.items() if key != "lng_daily_accounts"}
                                       for name, report in data["natural_gas_validation"].items()},
            "source_sha256": data["tracked"]}


def texts(data, language):
    c, t, n = data["cost"], data["time"], data["network"]
    s, runs = data["sensitivity"], data["runs"]
    x, g, b, d = "CCU-GH-FT", "GTL", "GTL-BH", "DAC-BH-FT"
    xa, xb, ga, gb = runs[x, "baseline"], runs[x, "summer_fixed"], runs["GTL-GH", "baseline"], runs["GTL-GH", "summer_fixed"]
    cx, cg, cb, tx, tg, tb, nx, ng, nb, nd = c.loc[x], c.loc[g], c.loc[b], t.loc[x], t.loc[g], t.loc[b], n.loc[x], n.loc[g], n.loc[b], n.loc[d]
    m = data["match25"]
    mean_dominant = [name for name in [b, g, x] if t.loc[name, "ChronicPenaltyPct"] >= t.loc[name, "RobustnessPenaltyPct"]]
    worst_dominant = [name for name in [b, g, x] if t.loc[name, "RobustnessPenaltyPct"] > t.loc[name, "ChronicPenaltyPct"]]
    storage = data["carbon_records"][b]["by_stage"]["storage_emissions"] / float(data["carbon_records"][b]["total_production_kg"]) / float(s.loc[b, "fuel_energy_mj_kg"]) * 1000
    storage_share = storage / cb.carbon_g_mj * 100
    totals = ", ".join(fmt(t.loc[name, "TotalPenaltyPct"], 1) + "%" for name in [b, g, x])
    total_cn = "、".join(fmt(t.loc[name, "TotalPenaltyPct"], 1) + "%" for name in [b, g, x])
    cost_delta = abs(cx.cost_cny_kg - cg.cost_cny_kg)
    ratio = cx.cost_cny_kg / cg.cost_cny_kg
    gap_max = max(row["gap"] for row in runs.values()) * 100
    retention_direction = ("rises" if xb["worst_week_retention"] >= xa["worst_week_retention"] else "falls",
                           "提高" if xb["worst_week_retention"] >= xa["worst_week_retention"] else "降低")
    mean_direction = ("rises" if xb["mean_utilization"] >= xa["mean_utilization"] else "falls",
                      "提高" if xb["mean_utilization"] >= xa["mean_utilization"] else "降低")
    gh_direction = ("rises" if gb["worst_week_retention"] >= ga["worst_week_retention"] else "falls",
                    "提高" if gb["worst_week_retention"] >= ga["worst_week_retention"] else "降低")
    demand_b = {round(runs[name, "summer_fixed"]["demand_kg"], 3) for name in ["CCU-GH-FT", "CCU-GH-MTJ", "DAC-GH-FT", "DAC-GH-MTJ", "GTL-GH", "CTL"]}
    if len(demand_b) != 1:
        raise ValueError("The fixed-asset cases do not share the same alternative airport demand")
    demand_million = next(iter(demand_b)) / 1e6
    out = {}
    if language == "en":
        dominance = []
        if worst_dominant:
            dominance.append(f"The additional worst-week component is larger for {names(worst_dominant)}")
        if mean_dominant:
            dominance.append(f"the mean-utilization component is larger for {names(mean_dominant)}")
        dominance_sentence = "; ".join(dominance) + "."
        match_phrase = f"capacity matched within 25 km represents {fmt(m['source_share_percent'], 1)}% and {fmt(m['target_share_percent'], 1)}% of their respective installed synthesis capacities"
        out[8] = f"Climate mitigation scenarios that treat sustainable aviation fuel as a unified category may overlook whether infrastructure deployed for near-term pathways transfers to later production configurations. We examine this question by optimizing full supply chains for 13 carbon–hydrogen–route combinations serving Beijing and Tianjin airports over four representative weeks. The cost–emission comparison reveals pronounced cost differences: natural-gas-to-liquids costs {fmt(cg.cost_cny_kg)} yuan/kg at {fmt(cg.fulfillment_ratio * 100)}% demand coverage, compared with {fmt(cx.cost_cny_kg)} yuan/kg and {fmt(cx.fulfillment_ratio * 100)}% coverage for captured industrial carbon and green hydrogen converted through Fischer–Tropsch synthesis. Spatially, the demand-driven network of the fossil pathway and the resource-driven network of the green-hydrogen pathway differ upstream but show partial synthesis-site overlap; {match_phrase}. Operationally, synthesis-capacity premiums reflect different combinations of mean utilization and worst-week loading, while the green-hydrogen pathway additionally exhibits low hydrogen-production utilization. These linked differences in cost, geography and operations show that synthetic aviation-fuel pathway switching can involve industrial regime reconstruction, even where some synthesis locations are shared. Net-zero scenarios therefore need to account explicitly for pathway-specific infrastructure transitions to avoid overstating how near-term deployment supports later aviation decarbonization."
        out[15] = f"The comparison reveals linked differences in cost, space and operations. In cost, GTL supplies {fmt(cg.fulfillment_ratio * 100)}% of demand at {fmt(cg.cost_cny_kg)} yuan/kg, whereas CCU-GH-FT supplies {fmt(cx.fulfillment_ratio * 100)}% at {fmt(cx.cost_cny_kg)} yuan/kg. In space, the {integer(ng['节点数量'])}-identifier demand-driven network of GTL and the {integer(nx['节点数量'])}-identifier resource-driven network of CCU-GH-FT differ upstream but show partial synthesis-site overlap: {match_phrase}. In operations, their mean-utilization and additional worst-week capacity premiums are {fmt(tg.ChronicPenaltyPct)}% and {fmt(tg.RobustnessPenaltyPct)}% for GTL, and {fmt(tx.ChronicPenaltyPct)}% and {fmt(tx.RobustnessPenaltyPct)}% for CCU-GH-FT; CCU-GH-FT hydrogen-production utilization is {fmt(tx.H2Mean)}%. These linked differences in cost, space and operations show that pathway switching can involve industrial regime reconstruction despite partial spatial continuity. This finding exposes the risk that net-zero scenarios overestimate the contribution of synthetic aviation fuels when they assume continuity between pathways with different infrastructure requirements, and provides empirical grounding for pathway-sensitive aviation decarbonization policy."
        out[17] = "The cost–emission comparison reveals distinct production architectures"
        out[18] = f"Within the cost–emission space defined by 13 synthetic aviation fuel pathways, the nondominated baseline configurations are {names(data['frontier'])} (Figure 1a). GTL-BH, GTL and CCU-GH-FT provide three representative production systems for the stage, spatial and operating comparisons. They differ in carbon–hydrogen sourcing and in the supporting process modules required to produce the same fuel. Their separation therefore spans discrete changes in production architecture. Each point describes costs and emissions at its own supplied quantity, so the frontier represents the finite set of modelled configurations; stage decomposition reveals the infrastructure changes underlying their positions."
        out[19] = f"The three representative systems show a pronounced cost difference at the captured-carbon–green-hydrogen combination. GTL-BH and GTL cost {fmt(cb.cost_cny_kg)} and {fmt(cg.cost_cny_kg)} CNY/kg (approximately {fmt(cb.cost_cny_kg / 7.1429)} and {fmt(cg.cost_cny_kg / 7.1429)} USD/kg), respectively, at demand coverage of {fmt(cb.fulfillment_ratio * 100)}% and {fmt(cg.fulfillment_ratio * 100)}%. CCU-GH-FT costs {fmt(cx.cost_cny_kg)} CNY/kg (approximately {fmt(cx.cost_cny_kg / 7.1429)} USD/kg), or {fmt(ratio)} times the GTL unit cost, while supplying {fmt(cx.fulfillment_ratio * 100)}% of demand (Figure 1a,b). The {fmt(cost_delta)} CNY/kg difference between GTL and CCU-GH-FT accompanies a change in supplied quantity. Their different carbon and hydrogen sources and associated facilities indicate that this cost difference involves structural changes between production systems."
        out[20] = f"The cost difference is associated with a reorganization of production architecture and its emissions structure. In captured-carbon pathways using green hydrogen, renewable electricity, hydrogen production and carbon capture become central inputs. The supply-chain carbon intensities of CCU-GH-FT and DAC-GH-FT are {fmt(cx.carbon_g_mj)} and {fmt(c.loc['DAC-GH-FT', 'carbon_g_mj'])} g CO₂e/MJ, respectively, while those of CCU-GH-MTJ and DAC-GH-MTJ are {fmt(c.loc['CCU-GH-MTJ', 'carbon_g_mj'])} and {fmt(c.loc['DAC-GH-MTJ', 'carbon_g_mj'])} g CO₂e/MJ. GTL-BH and GTL have supply-chain carbon intensities of {fmt(cb.carbon_g_mj)} and {fmt(cg.carbon_g_mj)} g CO₂e/MJ, respectively. Stage decomposition connects this cost–emission pattern to changes in production-system inputs and stages (Figure 1c–e)."
        out[21] = f"Stage decomposition reveals that the three representative pathways differ in the constituent elements of their production systems (Figure 1c–e). GTL and GTL-BH retain recognizable continuity through fossil-linked feedstock supply and fuel conversion. GTL-BH storage contributes {fmt(storage)} g CO₂e/MJ, or {fmt(storage_share)}% of its supply-chain intensity. CCU-GH-FT combines industrial carbon capture, electrolytic hydrogen and fuel synthesis, bringing electricity supply and capture facilities into the cost–emission structure. The different combinations of feedstock, hydrogen, capture and synthesis facilities distribute cost burdens across different stages and call for different sources of cost reduction. Figures S28–S29 show the cost and positive-emission compositions, with capture credits identified separately. Figure S25 extends the comparison to all thirteen pathways, and Figs. S36–S48 show learning projections for fixed production designs and supplied quantities."
        out[23] = f"Figure 1 Cost and supply-chain carbon accounts for thirteen synthetic aviation-fuel configurations. (a) Box tops and bottoms show 2026 baseline costs and 2036 fixed-design learning projections; internal marks identify intervening years. The horizontal axis shows baseline supply-chain carbon intensity, excluding aircraft combustion. (b) Baseline and projected cost differences from the assumed 8 CNY/kg ceiling; the shaded band represents 6–8 CNY/kg. (c–e) Polar bars show projected category costs. Outer rings show positive baseline supply-chain emissions before capture credits, in g CO₂e/MJ. GTL-BH, GTL and CCU-GH-FT supply {fmt(cb.fulfillment_ratio * 100)}%, {fmt(cg.fulfillment_ratio * 100)}% and {fmt(cx.fulfillment_ratio * 100)}% of demand, respectively; CTL is a partial-supply diagnostic. Costs exclude shortage penalties and use actual output. The envelope describes configurations at their respective service levels. Table S21 specifies service fractions and carbon-accounting variants; Supplementary Sections 6.6 and 7.8 give the learning assumptions and currency basis, respectively."
        out[25] = f"The production structure of CCU-GH-FT brings renewable electricity, electrolytic hydrogen and industrial carbon capture into the core requirements for fuel synthesis. Its baseline design installs {xa['h2_capacity_kg_h']:,.0f} kg/h of electrolyzer capacity and {xa['saf_capacity_kg_h']:,.0f} kg/h of fuel-synthesis capacity, producing {fmt(xa['production_kg'] / 1e6, 3)} million kg during the four weeks (Figure 1c–e). Electricity availability, hydrogen capacity, carbon supply and synthesis scale must consequently be configured together. The change in cost structure is thus accompanied by a change in the facilities that must be built and coordinated."
        out[27] = "Representative pathways exhibit distinct supply-chain spatial configurations"
        out[28] = "The differences among representative pathways extend from production structure to the spatial organization of physical infrastructure. Figure 2 compares the GTL, GTL-BH and CCU-GH-FT networks, with DAC-BH-FT providing a further spatial comparator. Their feedstock sourcing, facility distributions and transport connections follow different organizing principles. At the same time, GTL and CCU-GH-FT retain some synthesis capacity at common demand locations. These configurations connect distinct upstream networks with varying degrees of synthesis-site correspondence, making both dimensions relevant to infrastructure inheritability."
        out[29] = f"GTL-BH and GTL organize production and delivery differently. GTL-BH comprises {integer(nb['节点数量'])} positive-flow network identifiers and {integer(nb.saf_physical_sites)} physical synthesis sites. Its cross-regional logistics ratio is {fmt(nb['跨区域物流比例'])}%, and its quantity-weighted fuel-transport distance is {fmt(nb.SAF_distance_km)} km (Figure 2c,g,k). GTL comprises {integer(ng['节点数量'])} network identifiers and {integer(ng.saf_physical_sites)} physical synthesis sites, with a cross-regional logistics ratio of {fmt(ng['跨区域物流比例'])}% and mean fuel-transport distance of {fmt(ng.SAF_distance_km)} km (Figure 2b,f,j). GTL places synthesis at the two airport coordinates and meets {fmt(cg.fulfillment_ratio * 100)}% of demand at {fmt(cg.cost_cny_kg)} CNY/kg. These configurations show how feedstock and process choices change the network supporting the same airport demand. Scenario-level spatial and network indicators for all 13 pathways are summarized in Table S18, and pathway-specific transport-route maps are provided in the Supplementary Archive (Figs. S49–S63)."
        out[30] = f"The spatial expansion of CCU-GH-FT reorganizes hydrogen production and synthesis around dispersed renewable resources. Compared with GTL, which organizes synthesis at two airport sites, this configuration distributes hydrogen production and synthesis across multiple locations: it contains {integer(nx['节点数量'])} positive-flow network identifiers, {integer(nx.h2_supply_module_count)} hydrogen-production modules and {integer(nx.saf_module_count)} synthesis modules at {integer(nx.saf_physical_sites)} physical coordinates (Figure 2a,e,i). Production at these locations is connected through commodity flows with mean hydrogen and fuel transport distances of {fmt(nx.H2_distance_km)} and {fmt(nx.SAF_distance_km)} km, respectively; {fmt(nx['跨区域物流比例'])}% of aggregate mass flow travels more than 100 km. The network expansion thus centres on reorganizing facility locations and transport connections around resource sourcing. With {xa['saf_capacity_kg_h']:,.0f} kg/h of synthesis capacity, this configuration meets {fmt(cx.fulfillment_ratio * 100)}% of demand while retaining some synthesis capacity at airport demand sites, linking a dispersed upstream resource network to continuing demand-side locations. Resource distribution and demand geography therefore jointly shape the spatial organization of this pathway (Table S18 and Figs. S27 and S57)."
        position = "belongs to" if d in data["frontier"] else "lies outside"
        h2_route = f"The quantity-weighted hydrogen-transport distance is {fmt(nd.H2_distance_km)} km." if nd.H2_weekly_quantity > 1e-5 else "No positive explicit hydrogen-route flow is recorded."
        out[31] = f"The DAC-BH-FT configuration provides a different combination of spatial dispersion and cost–emission performance. It {position} the nondominated set in Figure 1a and contains {integer(nd['节点数量'])} positive-flow network identifiers and {integer(nd.saf_module_count)} synthesis modules at {integer(nd.saf_physical_sites)} physical coordinates. Its {integer(nd.h2_supply_module_count)} hydrogen-supply modules represent conditioning of industrial by-product hydrogen. {h2_route} The quantity-weighted fuel-transport distance is {fmt(nd.SAF_distance_km)} km, and {fmt(nd['跨区域物流比例'])}% of aggregate hydrogen, CO₂ and fuel mass flow travels more than 100 km (Table S18; Figure 2d,h,l). Facility dispersion, transport requirements and performance at the achieved service level therefore need to be considered together. The full DAC-BH-FT transport-route map is provided in the Supplementary Archive (Fig. S63)."
        out[33] = f"Figure 2 Spatial configurations and facility structures of representative synthetic aviation-fuel supply chains. (a–d) Supply-chain maps for CCU-GH-FT, GTL, GTL-BH and DAC-BH-FT, respectively. Their demand coverage is {fmt(cx.fulfillment_ratio * 100)}%, {fmt(cg.fulfillment_ratio * 100)}%, {fmt(cb.fulfillment_ratio * 100)}% and {fmt(c.loc[d, 'fulfillment_ratio'] * 100)}%, respectively. Symbols identify facility categories and line colours identify transported commodities. Endpoint-only transport records are shown as schematic connections; distance metrics use the recorded model route distances. (e–h) Network descriptors calculated under the definitions in Table S18, using positive flows and separately identified commodity units. (i–l) Facility-category counts, distinguishing network identifiers, equipment modules and co-located physical sites. Distance-threshold matching of synthesis capacity is reported in Table S24."
        threshold = "The same matched volume is obtained at thresholds from 1 to 100 km (Table S24)." if data["matching_constant"] else "Table S24 reports how the matched capacity changes across thresholds from 1 to 100 km."
        out[35] = f"Partial continuity at synthesis sites coexists with substantial upstream network reorganization. GTL and CCU-GH-FT organize hydrogen and carbon supply differently while locating some synthesis capacity at the same airport demand sites. After aggregating synthesis modules by coordinate, matching the {integer(ng.saf_physical_sites)} GTL synthesis sites against the {integer(nx.saf_physical_sites)} distinct CCU-GH-FT synthesis coordinates within 25 km gives a maximum one-to-one matched capacity of {m['matched_capacity_kg_h']:,.2f} kg/h, equal to {fmt(m['source_share_percent'])}% of GTL synthesis capacity and {fmt(m['target_share_percent'])}% of CCU-GH-FT capacity. {threshold} The shared airport locations therefore provide a geographic basis for connecting production to the same downstream demand, while hydrogen production, carbon sourcing and transport connections differ between the pathways. Infrastructure inheritability is consequently specific to supply-chain functions: common synthesis locations offer continuity alongside different upstream infrastructure requirements. The capacity match measures spatial correspondence, with equipment reuse subject to the engineering conditions discussed in Supplementary Section 7.9. A comparison of the representative network topologies is provided in the Supplementary Archive (Fig. S50)."
        out[36] = "Pathways differ in capacity utilization and worst-week loading requirements"
        out[38] = f"Figure 3 Temporal utilization and capacity requirements of three representative synthetic aviation-fuel configurations. (a) Mirrored utilization distributions for GTL-BH, GTL and CCU-GH-FT. Utilization increases from 0% at the centre to 100% at either outer edge, with hydrogen supply (H₂) on the left and fuel-synthesis plants on the right. Hydrogen denominators are active by-product-hydrogen conditioning capacity for GTL-BH and electrolyzer capacity for CCU-GH-FT (Supplementary Section 6.5.1). Ridge height represents a density estimate normalized to its own maximum for each distribution. Black and white dots indicate means and medians, respectively. No H₂ distribution is shown for GTL because this pathway does not rely on external hydrogen production. (b) Decomposition of capacity premium relative to full utilization into mean-utilization capacity premium (light blue) and additional worst-week capacity premium (dark blue). Dots indicate total premiums of {totals}, respectively; the percentages within the bars give the worst-week component as a share of each total. (c) Positions in capacity-premium space: mean-utilization premium on the x-axis, additional worst-week premium on the y-axis and total premium represented by bubble size. The dashed line y = x indicates equal contributions. (d) Descriptive utilization, low-load and renewable-availability metrics with cross-pathway normalization. Low-load episodes reset at each selected week boundary because the selected calendar weeks are nonconsecutive; long episodes last more than 12 hours. Demand coverage is {fmt(cb.fulfillment_ratio * 100)}%, {fmt(cg.fulfillment_ratio * 100)}% and {fmt(cx.fulfillment_ratio * 100)}% for GTL-BH, GTL and CCU-GH-FT, respectively."
        out[40] = f"Incorporating utilization loss and worst-week output contraction reveals divergence among the three representative pathways in the temporal-operational dimension. The divergence concerns both the concentration of low output across weeks and the production stages at which capacity is underutilized. Total synthesis-capacity premiums are {totals} for GTL-BH, GTL and CCU-GH-FT, respectively (Figure 3b,c). {dominance_sentence} CCU-GH-FT additionally has hydrogen-production utilization of {fmt(tx.H2Mean)}%. The decomposition expresses the extra capacity associated with producing a given quantity as a mean-utilization component and an additional worst-week component. The comparison thus links the scale of capacity requirements to the temporal and upstream–downstream organization of production. Formal definitions are provided in Methods and Supplementary Section 6.5.2; Fig. S32 shows the three-pathway decomposition and Table S23 reports all 13 pathways."
        out[41] = f"GTL-BH has mean fuel-synthesis utilization of {fmt(tb.SAFMean)}% and mean hydrogen-conditioning utilization of {fmt(tb.H2Mean)}% (Figure 3a). Its mean-utilization and additional worst-week capacity premiums are {fmt(tb.ChronicPenaltyPct)}% and {fmt(tb.RobustnessPenaltyPct)}%, respectively, and worst-week output retention is {fmt(tb.WorstWeekRetention)}% (Figure 3b,c). There are {integer(tb.LowLoadEpisodes12h)} low-load episodes lasting more than 12 h, with a maximum duration of {fmt(tb.MaxLowRunHours, 0)} h (Figure 3d). These measures connect average plant loading, contraction across the selected weeks and the duration of low-output operation."
        out[42] = f"GTL’s average utilization coexists with prolonged low output concentrated in its worst representative week. It maintains fuel-synthesis mean utilization of {fmt(tg.SAFMean)}% (Figure 3a), with an additional worst-week capacity premium of {fmt(tg.RobustnessPenaltyPct)}% and a mean-utilization capacity premium of {fmt(tg.ChronicPenaltyPct)}%. The worst-week component accounts for {fmt(tg.RobustnessSharePct)}% of the total premium of {fmt(tg.TotalPenaltyPct)}% (Figure 3b,c), and worst-week retention is {fmt(tg.WorstWeekRetention)}%. Across the four weeks, {integer(tg.LowLoadEpisodes12h)} low-load episode lasts more than 12 h; the maximum duration is {fmt(tg.MaxLowRunHours, 0)} h (Figure 3d). GTL meets each week’s demand, so prolonged low output reflects demand-following dispatch in the lowest-demand week. The distribution of production across weeks therefore has a central role in characterizing its operating requirements. All-pathway utilization distributions grouped by feedstock category are shown in Fig. S30; the four-panel overview in Fig. S31 reports worst-week retention, the two capacity-premium components and low-load episode statistics for all 13 pathways."
        out[43] = f"CCU-GH-FT extends the operating comparison from week-to-week synthesis output to coordination between hydrogen production and fuel synthesis. Mean H₂ utilization is {fmt(tx.H2Mean)}% and mean fuel-synthesis utilization is {fmt(tx.SAFMean)}% (Figure 3a). Its mean-utilization and additional worst-week capacity premiums are {fmt(tx.ChronicPenaltyPct)}% and {fmt(tx.RobustnessPenaltyPct)}%, respectively (Figure 3b,c), and worst-week retention is {fmt(tx.WorstWeekRetention)}%, compared with {fmt(tg.WorstWeekRetention)}% for GTL. Along the renewable-power-to-electrolysis-to-fuel-synthesis chain, a lightly utilized hydrogen-production stage supports a more heavily utilized synthesis stage. At this configuration’s demand fulfilment of {fmt(cx.fulfillment_ratio * 100)}%, these features create simultaneous requirements to coordinate the hydrogen and synthesis capacity mix and to allocate production across weeks. Its operating requirements consequently encompass both upstream capacity use and the timing of fuel production."
        out[44] = f"The fixed-asset comparison shows how mean utilization and relative worst-week performance respond to alternative resource–demand conditions. For weeks 1, 14, 27 and 44, CCU-GH-FT fuel-synthesis mean utilization {mean_direction[0]} from {fmt(xa['mean_utilization'] * 100)}% to {fmt(xb['mean_utilization'] * 100)}%, while worst-week retention {retention_direction[0]} from {fmt(xa['worst_week_retention'] * 100)}% to {fmt(xb['worst_week_retention'] * 100)}%. GTL-GH retention {gh_direction[0]} from {fmt(ga['worst_week_retention'] * 100)}% to {fmt(gb['worst_week_retention'] * 100)}%. The comparison fixes baseline synthesis and electrolyzer construction and capacity, together with DAC construction decisions where applicable; alternative-week demand totals {fmt(demand_million, 3)} million kg. Thus, even with production assets held fixed, weekly loading patterns change with the resource–demand combination. Table S32 reports all twelve A/B solutions and their supply levels, including the CTL partial-supply diagnostic with its original 90% minimum-service inequality removed in both week sets. Tables S27–S29 document all eighteen new solutions and their validation scope; Table S25 separately reports 52 leave-one-week-out calculations retaining the remaining stored schedules without re-optimization. Assessment of these configurations therefore needs to consider both average capacity use and relative worst-week performance."
        out[47] = f"Our analysis shows that synthetic aviation-fuel pathway switching can involve reorganization of supply-chain architecture across cost structure, spatial configuration and temporal-operational requirements. Pathways that appear adjacent within a shared SAF category do not necessarily share compatible upstream resource bases, network topologies or operating logics. In the Beijing–Tianjin case, GTL organizes synthesis at airport locations and meets {fmt(cg.fulfillment_ratio * 100)}% of demand, whereas CCU-GH-FT connects distributed upstream resources to synthesis capacity partly located at the same demand sites. Under the stated resource conditions, CCU-GH-FT meets {fmt(cx.fulfillment_ratio * 100)}% of demand and combines hydrogen-production utilization of {fmt(tx.H2Mean)}% with an additional worst-week synthesis-capacity premium of {fmt(tx.RobustnessPenaltyPct)}%. The coexistence of shared synthesis locations and different upstream requirements makes infrastructure inheritability a question about particular production and supply functions. Early low-cost capacity can become exposed to stranding if its resource connections and production equipment cannot support later pathways."
        sc = s.loc[x]
        out[65] = f"Combustion emissions and capture attribution affect the carbon-intensity interpretation of CCU-GH-FT. Tables S21–S22 add combustion emissions of {fmt(sc.combustion_addition_g_mj, 3)} g CO₂/MJ using the model’s {fmt(sc.fuel_carbon_mass_fraction)} fuel-carbon fraction and {fmt(sc.fuel_energy_mj_kg)} MJ/kg lower heating value. The baseline account is {fmt(sc.baseline_g_mj, 3)} g CO₂e/MJ; retaining the configured capture attribution and adding combustion gives {fmt(sc.with_combustion_g_mj, 3)} g CO₂e/MJ. Adding combustion without point-source capture attribution gives {fmt(sc.with_combustion_without_capture_attribution_g_mj, 3)} g CO₂e/MJ. The accounting distinguishes carbon recycling from permanent removal; changes in the source process require separate assessment. Comparison with the 89 g CO₂e/MJ conventional-fuel reference requires a consistent accounting boundary and explicit attribution of captured carbon."
        out[68] = f"The baseline CTL case cannot satisfy its 90% minimum-production requirement under the available resources: the necessary resource bound is {fmt(data['ctl_upper'] * 100, 4)}%. A partial-supply diagnostic removes that inequality while retaining all other constraints; its production assets are fixed for the alternative-week test, which uses the same diagnostic service scope."
        out[81] = f"The baseline CTL model cannot satisfy its 90% minimum-production requirement; its necessary resource upper bound is {fmt(data['ctl_upper'] * 100, 4)}%. The partial-supply diagnostic removes this inequality while retaining the remaining constraints, and its baseline production assets define the alternative-week comparison under the same diagnostic service scope."
        out[86] = f"The optimization model was implemented in Python. Twelve configurations were recomputed with COPT 8.0.7: twelve baseline solutions and six fixed-production-design alternative-week solutions. All eighteen solutions meet a 1% total-objective-gap criterion; the largest achieved gap is {fmt(gap_max, 4)}%. Feasible schedules were checked against every row, bound and integer decision of the stated model scope, with global lower bounds obtained from valid routing projections. Independent checks cover supply, electricity, hydrogen, CO₂ where applicable, fuel, inventory, transport and demand; all six fixed-asset comparisons pass. CTL A and B use the partial-supply diagnostic scope described above. The gap applies to the total objective, which includes shortage penalties, and does not bound the relative error in unit production cost. The original GTL comparator retains its archived Gurobi 12.0.3 result. Tables S27–S29 report provenance, bounds and validation scope. "
    else:
        dominance = []
        if worst_dominant:
            dominance.append(f"{names(worst_dominant, 'zh')}的最差周附加分量较大")
        if mean_dominant:
            dominance.append(f"{names(mean_dominant, 'zh')}的平均利用率分量较大")
        dominance_sentence = "；".join(dominance) + "。"
        out[8] = f"将可持续航空燃料作为统一类别的气候减缓情景，可能忽略近期路径所建设的基础设施能否支持后续生产配置。本研究对服务北京和天津机场的13种碳源—氢源—合成路线组合开展完整供应链优化，比较四个代表周内的运行。成本—排放结果表明，不同路径的成本差异明显：天然气制液体燃料的成本为{fmt(cg.cost_cny_kg)}元/kg，满足{fmt(cg.fulfillment_ratio * 100)}%的需求；工业捕集碳与绿氢通过费托合成生产燃料的成本为{fmt(cx.cost_cny_kg)}元/kg，满足{fmt(cx.fulfillment_ratio * 100)}%的需求。空间上，需求驱动的化石路径与资源驱动的绿氢路径具有不同的上游网络，同时保留部分合成厂址重合；25 km阈值内的匹配产能分别占两者合成装机的{fmt(m['source_share_percent'], 1)}%和{fmt(m['target_share_percent'], 1)}%。运行上，合成产能溢价由平均利用率和最差周负荷的不同组合构成，绿氢路径还具有制氢端低利用率的特征。成本、地理与运行的关联差异表明，即使共享部分合成地点，合成航空燃料路径转换仍可能涉及工业体系重组。净零情景因此需要明确考虑各路径的基础设施转换条件，避免高估近期部署对后续航空脱碳的支持。"
        out[15] = f"比较结果显示，成本、空间与运行差异相互联系。在成本上，GTL以{fmt(cg.cost_cny_kg)}元/kg满足{fmt(cg.fulfillment_ratio * 100)}%的需求，CCU-GH-FT以{fmt(cx.cost_cny_kg)}元/kg满足{fmt(cx.fulfillment_ratio * 100)}%的需求。在空间上，GTL的需求驱动网络包含{integer(ng['节点数量'])}个网络节点标识，CCU-GH-FT的资源驱动网络包含{integer(nx['节点数量'])}个网络节点标识，两者上游配置不同，但保留部分合成厂址重合：25 km阈值内的匹配产能分别占两者合成装机的{fmt(m['source_share_percent'])}%和{fmt(m['target_share_percent'])}%。在运行上，GTL的平均利用率容量溢价与最差周附加容量溢价分别为{fmt(tg.ChronicPenaltyPct)}%和{fmt(tg.RobustnessPenaltyPct)}%，CCU-GH-FT分别为{fmt(tx.ChronicPenaltyPct)}%和{fmt(tx.RobustnessPenaltyPct)}%；后者制氢端利用率为{fmt(tx.H2Mean)}%。这些关联差异表明，即使存在部分空间连续性，路径转换仍可能涉及工业体系重组。当净零情景假定基础设施要求不同的路径之间具有连续性时，可能高估合成航空燃料的贡献；这一发现为按路径特征制定航空脱碳政策提供了实证基础。"
        out[17] = "成本—排放比较揭示不同的生产体系结构"
        out[18] = f"在13条合成航空燃料路径构成的成本—排放空间中，基准非支配配置为{names(data['frontier'], 'zh')}（图1a）。GTL-BH、GTL和CCU-GH-FT作为三种代表性生产体系，用于工序、空间与运行比较。它们的碳氢来源及生产同一种燃料所需的配套工艺模块不同，因此在生产结构上体现离散的改变。各数据点均表示自身实际供给量下的成本与排放，前沿描述有限的模型配置集合；工序分解进一步解释其位置背后的基础设施差异。"
        out[19] = f"三种代表性体系在工业捕集碳与绿氢组合处呈现明显成本差异。GTL-BH与GTL的单位成本分别为{fmt(cb.cost_cny_kg)}和{fmt(cg.cost_cny_kg)} CNY/kg，约合{fmt(cb.cost_cny_kg / 7.1429)}和{fmt(cg.cost_cny_kg / 7.1429)} USD/kg，需求满足率分别为{fmt(cb.fulfillment_ratio * 100)}%和{fmt(cg.fulfillment_ratio * 100)}%。CCU-GH-FT的成本为{fmt(cx.cost_cny_kg)} CNY/kg，约合{fmt(cx.cost_cny_kg / 7.1429)} USD/kg，为GTL的{fmt(ratio)}倍，同时满足{fmt(cx.fulfillment_ratio * 100)}%的需求（图1a、b）。GTL与CCU-GH-FT之间{fmt(cost_delta)} CNY/kg的成本差异伴随供给量变化。两者采用不同的碳氢来源及配套设施，说明成本变化涉及生产体系结构的改变。"
        out[20] = f"成本差异与生产结构及排放构成的重组相联系。在捕集碳与绿氢路径中，可再生电力、制氢和碳捕集成为核心投入。CCU-GH-FT与DAC-GH-FT的供应链碳强度分别为{fmt(cx.carbon_g_mj)}和{fmt(c.loc['DAC-GH-FT', 'carbon_g_mj'])} g CO₂e/MJ，CCU-GH-MTJ与DAC-GH-MTJ分别为{fmt(c.loc['CCU-GH-MTJ', 'carbon_g_mj'])}和{fmt(c.loc['DAC-GH-MTJ', 'carbon_g_mj'])} g CO₂e/MJ。GTL-BH与GTL的供应链碳强度分别为{fmt(cb.carbon_g_mj)}和{fmt(cg.carbon_g_mj)} g CO₂e/MJ。工序分解将这一成本—排放格局与生产体系投入及工序变化联系起来（图1c—e）。"
        out[21] = f"工序分解显示，三条代表性路径的生产体系构成不同（图1c—e）。GTL与GTL-BH在化石原料供应和燃料转化方面保留可识别的连续性。GTL-BH的储存排放为{fmt(storage)} g CO₂e/MJ，占其供应链碳强度的{fmt(storage_share)}%。CCU-GH-FT结合工业碳捕集、电解制氢和燃料合成，使电力供应与捕集设施进入成本—排放结构的核心。原料、供氢、捕集与合成设施的不同组合，将成本负担分配至不同工序，并对应不同的降本来源。图S28—S29展示成本与正排放构成，捕集抵扣单独列示。图S25扩展至全部13条路径，图S36—S48展示固定生产设计和供给量下的学习投影。"
        out[23] = f"图1 13种合成航空燃料配置的成本与供应链碳核算。（a）箱体顶部和底部分别表示2026年基准成本与2036年固定设计学习投影，内部标记表示中间年份。横轴为不含航空器燃烧的基准供应链碳强度。（b）基准及投影成本与假定8 CNY/kg价格上限的差额，阴影表示6—8 CNY/kg区间。（c—e）极坐标柱表示各类别投影成本，外环表示扣除捕集抵扣前的正供应链排放，单位为g CO₂e/MJ。GTL-BH、GTL和CCU-GH-FT分别满足{fmt(cb.fulfillment_ratio * 100)}%、{fmt(cg.fulfillment_ratio * 100)}%和{fmt(cx.fulfillment_ratio * 100)}%的需求；CTL为部分供给诊断情景。成本剔除短缺惩罚并按实际产量计算，包络线描述各自服务水平下的配置。表S21列示供给比例与碳核算变体，补充材料第6.6和7.8节分别说明学习假设与币值基础。"
        out[25] = f"CCU-GH-FT的生产结构将可再生电力、电解制氢与工业碳捕集纳入燃料合成的核心要求。其基准设计配置{xa['h2_capacity_kg_h']:,.0f} kg/h的电解制氢产能和{xa['saf_capacity_kg_h']:,.0f} kg/h的燃料合成产能，四个代表周产量为{fmt(xa['production_kg'] / 1e6, 3)}×10⁶ kg（图1c—e）。因此，需要联合配置电力可用量、制氢产能、碳供给与合成规模。成本结构的变化同时伴随所需建设与协调的设施变化。"
        out[27] = "代表性路径具有不同的供应链空间配置"
        out[28] = "代表性路径的差异从生产结构延伸至基础设施的空间组织。图2比较GTL、GTL-BH与CCU-GH-FT网络，并以DAC-BH-FT作为进一步的空间对照。原料来源、设施分布和运输连接遵循不同的组织原则。同时，GTL与CCU-GH-FT仍在共同需求地点保留部分合成产能。不同上游网络与不同程度的合成厂址对应关系相联系，使这两个维度共同影响基础设施沿用条件。"
        out[29] = f"GTL-BH与GTL采用不同的生产和交付组织。GTL-BH包含{integer(nb['节点数量'])}个正流量网络节点标识及{integer(nb.saf_physical_sites)}个实际合成厂址，跨区域物流比例为{fmt(nb['跨区域物流比例'])}%，按运输量加权的燃料运输距离为{fmt(nb.SAF_distance_km)} km（图2c、g、k）。GTL包含{integer(ng['节点数量'])}个网络节点标识及{integer(ng.saf_physical_sites)}个实际合成厂址，跨区域物流比例为{fmt(ng['跨区域物流比例'])}%，平均燃料运输距离为{fmt(ng.SAF_distance_km)} km（图2b、f、j）。GTL在两个机场坐标配置合成，以{fmt(cg.cost_cny_kg)} CNY/kg满足{fmt(cg.fulfillment_ratio * 100)}%的需求。这些配置展示了原料与工艺选择如何改变服务同一机场需求的网络。全部13条路径的空间和网络指标见表S18，各路径运输地图见补充图册图S49—S63。"
        out[30] = f"CCU-GH-FT围绕分散的可再生资源重新组织制氢与合成环节。与依托两个机场厂址的GTL相比，其制氢与合成采用多点布局：包含{integer(nx['节点数量'])}个正流量网络节点标识、{integer(nx.h2_supply_module_count)}个制氢模块，以及位于{integer(nx.saf_physical_sites)}个实际坐标处的{integer(nx.saf_module_count)}个合成模块（图2a、e、i）。各地点通过物料流连接，氢气与燃料平均运输距离分别为{fmt(nx.H2_distance_km)}和{fmt(nx.SAF_distance_km)} km，超过100 km的质量流量占{fmt(nx['跨区域物流比例'])}%。网络扩展表现为围绕资源获取改变设施布置和运输连接。该配置以{xa['saf_capacity_kg_h']:,.0f} kg/h的合成产能满足{fmt(cx.fulfillment_ratio * 100)}%的需求，部分合成产能仍位于机场需求地点，使分散的上游资源网络与延续的需求端厂址相联系。资源分布与需求位置共同塑造了该路径的空间组织（表S18及图S27、S57）。"
        position = "进入" if d in data["frontier"] else "未进入"
        h2_route = f"按运输量加权的氢气运输距离为{fmt(nd.H2_distance_km)} km。" if nd.H2_weekly_quantity > 1e-5 else "未记录显式氢气运输路线的正流量。"
        out[31] = f"DAC-BH-FT展示另一种空间分散与成本—排放表现的组合。该配置{position}图1a的非支配集合，包含{integer(nd['节点数量'])}个正流量网络节点标识，以及分布在{integer(nd.saf_physical_sites)}个实际坐标处的{integer(nd.saf_module_count)}个合成模块。其{integer(nd.h2_supply_module_count)}个供氢模块用于工业副产氢净化。{h2_route}按运输量加权的燃料运输距离为{fmt(nd.SAF_distance_km)} km，氢气、CO₂与燃料合计质量流量中，超过100 km的部分占{fmt(nd['跨区域物流比例'])}%（表S18；图2d、h、l）。因此，需要联合考虑设施分散、运输要求与实际供给水平下的表现。完整运输路线图见补充图册图S63。"
        out[33] = f"图2 代表性合成航空燃料供应链的空间配置与设施结构。（a—d）依次为CCU-GH-FT、GTL、GTL-BH和DAC-BH-FT的供应链地图，需求满足率依次为{fmt(cx.fulfillment_ratio * 100)}%、{fmt(cg.fulfillment_ratio * 100)}%、{fmt(cb.fulfillment_ratio * 100)}%和{fmt(c.loc[d, 'fulfillment_ratio'] * 100)}%。符号表示设施类别，线条颜色表示运输物料。仅保存起终点的运输记录采用示意连接，距离指标使用模型记录的路线距离。（e—h）按表S18定义计算的网络指标，使用正流量并区分不同物料单位。（i—l）设施类别数量，区分网络节点标识、设备模块与共址的实际厂址。距离阈值下的合成产能匹配见表S24。"
        threshold = "1—100 km范围内各阈值的匹配产能相同（表S24）。" if data["matching_constant"] else "表S24列示1—100 km范围内各阈值对应的匹配产能。"
        out[35] = f"合成厂址的局部连续性与上游供给网络的明显重组可以并存。GTL与CCU-GH-FT采用不同的供氢和供碳组织，同时在共同机场需求地点配置部分合成产能。按坐标合并合成模块后，GTL的{integer(ng.saf_physical_sites)}个合成厂址与CCU-GH-FT的{integer(nx.saf_physical_sites)}个合成坐标在25 km阈值内的最大一对一匹配产能为{m['matched_capacity_kg_h']:,.2f} kg/h，分别占GTL合成产能的{fmt(m['source_share_percent'])}%和CCU-GH-FT合成产能的{fmt(m['target_share_percent'])}%。{threshold}共同机场地点为两条路径连接同一下游需求提供了地理基础，同时两者的制氢、供碳和运输连接采用不同配置。基础设施沿用因此需要按供应链功能判断：共同合成地点提供连续性，上游基础设施具有不同要求。产能匹配衡量空间对应关系，设备复用取决于补充材料第7.9节讨论的工程条件。代表性网络拓扑比较见补充图册图S50。"
        out[36] = "不同路径的产能利用与最差周负荷要求存在差异"
        out[38] = f"图3 三种代表性合成航空燃料配置的时序利用率与产能要求。（a）GTL-BH、GTL和CCU-GH-FT的镜像利用率分布，从中央0%向两侧外缘100%递增，左侧为供氢，右侧为燃料合成。供氢分母分别采用GTL-BH的已建设副产氢净化产能和CCU-GH-FT的电解产能（补充材料第6.5.1节）。每条密度曲线按自身最大值归一化，黑白圆点分别表示均值和中位数。GTL不依赖外部制氢，未展示H₂分布。（b）相对于满负荷的产能溢价分解为平均利用率容量溢价（浅蓝色）与最差周附加容量溢价（深蓝色）。圆点表示依次为{total_cn}的总容量溢价，柱内百分比为最差周部分占总量的比例。（c）横轴为平均利用率容量溢价，纵轴为最差周附加容量溢价，气泡大小表示总容量溢价；虚线y=x表示两个分量相等。（d）经跨路径归一化的利用率、低负荷与可再生能源可用性描述指标。非连续代表周的低负荷事件在每周边界重新计数，长时事件指超过12小时的事件。GTL-BH、GTL和CCU-GH-FT的需求满足率依次为{fmt(cb.fulfillment_ratio * 100)}%、{fmt(cg.fulfillment_ratio * 100)}%和{fmt(cx.fulfillment_ratio * 100)}%。"
        out[40] = f"纳入利用率损失与最差周产出收缩后，三条代表性路径在时序运行维度呈现差异，既涉及低产出在各周之间的集中程度，也涉及产能利用不足发生的生产环节。GTL-BH、GTL和CCU-GH-FT的合成端总容量溢价依次为{total_cn}（图3b、c）。{dominance_sentence}CCU-GH-FT的制氢端利用率为{fmt(tx.H2Mean)}%。分解将完成给定产量所需的额外产能表达为平均利用率分量与最差周附加分量，从而联系产能要求的大小与生产在时间及上下游之间的组织。正式定义见方法和补充材料第6.5.2节，三条路径的分解见图S32，全部13条路径的结果见表S23。"
        out[41] = f"GTL-BH的燃料合成平均利用率为{fmt(tb.SAFMean)}%，供氢净化平均利用率为{fmt(tb.H2Mean)}%（图3a）。其平均利用率容量溢价与最差周附加容量溢价分别为{fmt(tb.ChronicPenaltyPct)}%和{fmt(tb.RobustnessPenaltyPct)}%，最差周产出保持率为{fmt(tb.WorstWeekRetention)}%（图3b、c）。超过12 h的低负荷事件为{integer(tb.LowLoadEpisodes12h)}次，最长持续{fmt(tb.MaxLowRunHours, 0)} h（图3d）。这些指标共同描述平均厂站负荷、各代表周之间的产出收缩和低产出运行持续时间。"
        out[42] = f"GTL的平均利用率与集中于最差代表周的长时低负荷并存。其燃料合成平均利用率为{fmt(tg.SAFMean)}%（图3a），最差周附加容量溢价为{fmt(tg.RobustnessPenaltyPct)}%，平均利用率容量溢价为{fmt(tg.ChronicPenaltyPct)}%。最差周分量占总容量溢价{fmt(tg.TotalPenaltyPct)}%的{fmt(tg.RobustnessSharePct)}%（图3b、c），最差周保持率为{fmt(tg.WorstWeekRetention)}%。四个代表周中，超过12 h的低负荷事件为{integer(tg.LowLoadEpisodes12h)}次，最长持续{fmt(tg.MaxLowRunHours, 0)} h（图3d）。GTL在各周均满足需求，因此长时低负荷体现最低需求周中的随需求调度，生产任务在各周之间的分布由此影响运行要求。全部路径按原料类别分组的利用率分布见图S30，图S31列示全部13条路径的最差周保持率、两类容量溢价及低负荷事件统计。"
        out[43] = f"CCU-GH-FT将运行比较从合成端周间产出延伸至制氢与燃料合成的协调。其H₂平均利用率为{fmt(tx.H2Mean)}%，燃料合成平均利用率为{fmt(tx.SAFMean)}%（图3a）；平均利用率容量溢价和最差周附加容量溢价分别为{fmt(tx.ChronicPenaltyPct)}%和{fmt(tx.RobustnessPenaltyPct)}%（图3b、c）。最差周保持率为{fmt(tx.WorstWeekRetention)}%，GTL为{fmt(tg.WorstWeekRetention)}%。沿可再生电力—电解制氢—燃料合成链条，较低利用率的制氢环节支撑较高利用率的合成环节。在{fmt(cx.fulfillment_ratio * 100)}%的需求满足率下，需要同时协调制氢与合成产能配比，并安排各周生产任务。因此，其运行要求同时涉及上游产能使用与燃料生产时间。"
        out[44] = f"固定资产比较显示平均利用率与最差周相对表现如何响应替代资源—需求条件。在第1、14、27和44周组合下，CCU-GH-FT合成平均利用率由{fmt(xa['mean_utilization'] * 100)}%{mean_direction[1]}至{fmt(xb['mean_utilization'] * 100)}%，最差周保持率由{fmt(xa['worst_week_retention'] * 100)}%{retention_direction[1]}至{fmt(xb['worst_week_retention'] * 100)}%。GTL-GH的保持率由{fmt(ga['worst_week_retention'] * 100)}%{gh_direction[1]}至{fmt(gb['worst_week_retention'] * 100)}%。比较固定基准合成与电解设施的建设和产能，并在适用情况下固定DAC建设；替代周需求为{fmt(demand_million, 3)}×10⁶ kg。因此，即使生产资产固定，各周负荷仍随资源—需求组合变化。表S32列示12个A/B解及其供给水平，其中CTL两个周组合均为移除原90%最低服务约束的部分供给诊断。表S27—S29记录全部18个新解及验证范围；表S25另列52组保留其余已存调度、不重新优化的逐次剔周计算。运行评价需要同时考虑平均产能使用与最差周相对表现。"
        out[47] = f"本研究表明，合成航空燃料路径转换可能涉及成本结构、空间配置与时序运行要求三个方面的供应链体系重组。同属SAF类别的相近路径，未必具有兼容的上游资源、网络结构与运行逻辑。北京—天津案例中，GTL在机场位置组织合成并满足{fmt(cg.fulfillment_ratio * 100)}%的需求，CCU-GH-FT将分散的上游资源与部分位于相同需求地点的合成产能连接。在给定资源条件下，CCU-GH-FT满足{fmt(cx.fulfillment_ratio * 100)}%的需求，制氢端利用率为{fmt(tx.H2Mean)}%，合成端最差周附加容量溢价为{fmt(tx.RobustnessPenaltyPct)}%。共享合成厂址与上游要求差异并存，使基础设施沿用需要按具体生产和供给功能判断。早期低成本产能的资源连接与设备如果无法支持后续路径，就可能面临搁浅风险。"
        sc = s.loc[x]
        out[65] = f"燃烧排放与捕集碳归属共同影响CCU-GH-FT的碳强度判断。表S21—S22根据{fmt(sc.fuel_carbon_mass_fraction)}的燃料碳质量分数和{fmt(sc.fuel_energy_mj_kg)} MJ/kg的低位热值，计入{fmt(sc.combustion_addition_g_mj, 3)} g CO₂/MJ的燃烧排放。基准核算为{fmt(sc.baseline_g_mj, 3)} g CO₂e/MJ，保留既定捕集归属并计入燃烧时为{fmt(sc.with_combustion_g_mj, 3)} g CO₂e/MJ；计入燃烧且不计点源捕集归属时为{fmt(sc.with_combustion_without_capture_attribution_g_mj, 3)} g CO₂e/MJ。核算区分碳循环利用与永久碳移除，来源工艺变化的影响需另外评价。与89 g CO₂e/MJ传统燃料参考值比较时，需要采用一致的核算边界，并明确捕集碳归属。"
        out[68] = f"基准CTL情景在现有资源条件下无法满足90%的最低产量要求，其必要资源供给上界为{fmt(data['ctl_upper'] * 100, 4)}%。部分供给诊断移除该项不等式并保留其余约束，固定该诊断的生产资产进行替代周测试，两个周组合采用相同的诊断服务范围。"
        out[81] = f"基准CTL模型无法满足90%的最低产量要求，其必要资源供给上界为{fmt(data['ctl_upper'] * 100, 4)}%。部分供给诊断移除该项不等式并保留其余约束，固定其基准生产资产开展相同诊断服务范围下的替代周比较。"
        out[86] = f"优化模型采用Python实现，使用COPT 8.0.7重新计算12种配置，得到12个基准解及6个固定生产设计的替代周解。全部18个解满足总目标相对间隙不超过1%的标准，最大实际间隙为{fmt(gap_max, 4)}%。可行调度经过所述模型范围的全部约束、变量边界与整数决策检查，路线投影提供有效全局下界。独立验证覆盖供给、电力、氢、适用情景的CO₂、燃料、库存、运输与需求，6组固定资产比较均通过。CTL的A/B解采用前述部分供给诊断范围。间隙针对含短缺惩罚的总目标，不作为单位生产成本相对误差上限。原GTL对照保留归档的Gurobi 12.0.3结果。表S27—S29列示来源、上下界与验证范围。"
    if language == "en":
        out[29] = out[29].replace(f"Its cross-regional logistics ratio is {fmt(nb['跨区域物流比例'])}%, and its quantity-weighted fuel-transport distance is {fmt(nb.SAF_distance_km)} km",
                                f"Its {integer(nb.h2_supply_module_count)} hydrogen-conditioning modules support airport synthesis; the cross-regional logistics ratio is {fmt(nb['跨区域物流比例'])}%, with quantity-weighted hydrogen and fuel transport distances of {fmt(nb.H2_distance_km)} and {fmt(nb.SAF_distance_km)} km, respectively")
        if len(data["directly_shared_synthesis_sites"]) == 1:
            out[28] = out[28].replace("at common demand locations", "at a common demand location")
            out[30] = out[30].replace("at airport demand sites", "at an airport demand site")
            out[35] = out[35].replace("at the same airport demand sites", "at a shared airport demand site").replace("The shared airport locations therefore provide", "The shared airport site therefore provides")
            out[47] = out[47].replace("partly located at the same demand sites", "partly located at a shared demand site")
            out[51] = "A shared demand-side synthesis site provides"
        else:
            out[51] = "Shared demand-side synthesis sites provide"
    else:
        out[29] = out[29].replace(f"跨区域物流比例为{fmt(nb['跨区域物流比例'])}%，按运输量加权的燃料运输距离为{fmt(nb.SAF_distance_km)} km",
                                f"配置{integer(nb.h2_supply_module_count)}个供氢净化模块支持机场合成，跨区域物流比例为{fmt(nb['跨区域物流比例'])}%，按运输量加权的氢气与燃料运输距离分别为{fmt(nb.H2_distance_km)}和{fmt(nb.SAF_distance_km)} km")
        out[51] = "共同需求地点的合成厂址提供了空间上的衔接机会"
    return out


def local_segment(index, paragraph, language):
    if index == 51:
        anchor = "Common demand-side synthesis sites provide" if language == "en" else "共同需求地点的合成厂址提供了空间上的衔接机会"
        if paragraph.count(anchor) != 1:
            raise ValueError(f"The concluding shared-site phrase needs a unique local anchor: {language}")
        return anchor
    if index == 15:
        anchor = "The comparison reveals linked differences" if language == "en" else "比较结果"
        if anchor not in paragraph:
            anchor = "比较揭示" if language == "zh" else anchor
        if paragraph.count(anchor) != 1:
            raise ValueError(f"The existing Introduction result summary needs a unique local anchor: {language}")
        return paragraph[paragraph.index(anchor):]
    if index in [68, 81]:
        anchor = "The baseline CTL" if language == "en" else "基准CTL"
        if paragraph.count(anchor) != 1:
            raise ValueError(f"The current CTL scope sentence needs a unique local anchor: {index}/{language}")
        end_anchor = "Result provenance" if language == "en" and index == 81 else "各情景的结果来源" if language == "zh" and index == 81 else None
        end = paragraph.index(end_anchor) if end_anchor else len(paragraph)
        return paragraph[paragraph.index(anchor):end].rstrip()
    if index == 86:
        anchor = "USD equivalents" if language == "en" else "美元等值"
        if paragraph.count(anchor) != 1:
            raise ValueError(f"The existing solver paragraph needs a unique currency boundary: {language}")
        return paragraph[:paragraph.index(anchor)]
    return paragraph


def prepare(locations, data):
    documents = {}
    for name, bound in locations["documents"].items():
        language = bound["language"]
        replacement = texts(data, language)
        edits = []
        for index in TARGETS:
            before_paragraph = bound["paragraphs"][str(index)]
            before = local_segment(index, before_paragraph, language)
            after = replacement[index]
            if before == after:
                continue
            if before_paragraph.count(before) != 1:
                raise ValueError(f"Nonunique result phrase: {name}/P{index}")
            edits.append({"paragraph_index": bound["paragraph_indices"][str(index)], "clean_reference_index": index, "before": before, "after": after,
                          "paragraph_before_sha256": hashlib.sha256(before_paragraph.encode()).hexdigest(),
                          "paragraph_after": before_paragraph.replace(before, after, 1)})
        documents[name] = {"source_path": bound["source_path"], "source_sha256": bound["source_sha256"], "language": language, "edits": edits}
    return {"schema_version": 1, "ready": True, "docx_saved": False, "accepted_new_runs": 18,
            "source_sha256": data["tracked"], "actual_nondominated_baselines": data["frontier"],
            "representative_pathways": ["GTL-BH", "GTL", "CCU-GH-FT"],
            "scope": "Necessary result facts, headings and captions; acknowledgments and availability statements are outside this plan.",
            "documents": documents}


def apply_replacements(document, plan, filename, marked=False):
    if plan.get("ready") is not True or plan.get("accepted_new_runs") != 18:
        raise ValueError("Result replacements require the complete accepted numerical plan")
    for path, expected in plan["source_sha256"].items():
        if digest(Path(path)) != expected:
            raise ValueError(f"Prepared result source changed: {path}")
    for edit in plan["documents"][filename]["edits"]:
        paragraph = document.paragraphs[edit["paragraph_index"]]
        current = accepted(paragraph)
        if hashlib.sha256(current.encode()).hexdigest() != edit["paragraph_before_sha256"]:
            raise ValueError(f"Current paragraph changed: {filename}/P{edit['paragraph_index']}")
        replace_phrase(paragraph, edit["before"], edit["after"], marked=marked)
        if accepted(paragraph) != edit["paragraph_after"]:
            raise ValueError(f"Accepted replacement differs: {filename}/P{edit['paragraph_index']}")
    return document


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["inspect", "prepare"])
    parser.add_argument("--packet", type=Path, default=PACKET)
    parser.add_argument("--locations", type=Path, default=HERE / "main_result_locations.json")
    parser.add_argument("--out", type=Path, default=HERE / "main_result_replacements.json")
    parser.add_argument("--values-out", type=Path, default=HERE / "publication_values.json")
    parser.add_argument("--results", type=Path, default=WORK / "results")
    parser.add_argument("--accounts", type=Path, default=WORK / "postprocess/accounts")
    parser.add_argument("--result-tables", type=Path, default=WORK / "result_tables/output")
    parser.add_argument("--costs", type=Path, default=WORK / "figures_cost/adapted_data")
    parser.add_argument("--spatial", type=Path, default=WORK / "figures_spatial/network_data")
    parser.add_argument("--temporal", type=Path, default=WORK / "figures_temporal")
    args = parser.parse_args()
    locations = inspect(args.packet)
    if args.mode == "inspect":
        write_json(args.locations, locations)
        print(json.dumps({"mode": "inspect", "documents": len(locations["documents"]), "paragraphs_per_document": len(TARGETS), "docx_saved": False}))
    else:
        data = facts(args.results, args.accounts, args.result_tables, args.costs, args.spatial, args.temporal)
        plan = prepare(locations, data)
        values = publication_values(data)
        write_json(args.values_out, values)
        plan["publication_values_path"] = str(args.values_out.resolve())
        plan["publication_values_sha256"] = digest(args.values_out)
        write_json(args.out, plan)
        print(json.dumps({"mode": "prepare", "documents": len(plan["documents"]), "actual_nondominated_baselines": plan["actual_nondominated_baselines"], "docx_saved": False}))


if __name__ == "__main__":
    main()
