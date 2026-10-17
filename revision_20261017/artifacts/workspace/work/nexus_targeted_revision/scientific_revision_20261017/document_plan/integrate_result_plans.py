from pathlib import Path
import argparse
import json
import re

from docx import Document

from main_result_replacements import apply_replacements as apply_main
from si_result_replacements import apply_replacements as apply_si
from update_si_results import PACKET, apply_updates, digest, read_json, write_json
from word_edits import accepted


HERE = Path(__file__).resolve().parent
PAIRS = [("Manuscript_Clean.docx", "Manuscript_Marked_Changes.docx"),
         ("作者审阅/Manuscript_Chinese.docx", "作者审阅/Manuscript_Marked_Changes_Chinese.docx"),
         ("Supplementary_Information.docx", "Supplementary_Information_Marked_Changes.docx"),
         ("Supplementary_Figure_Archive.docx", "Supplementary_Figure_Archive_Marked_Changes.docx")]


def normalized_text(document):
    paragraphs = [accepted(p) for p in document.paragraphs if accepted(p).strip()
                  and not re.fullmatch(r"Revised Figure \d+|修改图\d+|\(\d+\)", accepted(p).strip())]
    return re.sub(r"\s+", "", "".join(paragraphs))


def accepted_tables(document):
    tables = [[["\n".join(accepted(p) for p in cell.paragraphs) for cell in row.cells] for row in table.rows]
              for table in document.tables]
    return [[row for row in table if any(cell.strip() for cell in row)] for table in tables
            if any(any(cell.strip() for cell in row) for row in table)]


def citations(document):
    return [re.sub(r"\s+", "", run.text) for paragraph in document.paragraphs for run in paragraph.runs
            if run.font.superscript and not run.font.strike and re.search(r"\d", run.text)]


def integrate(packet, main_plan, si_plan, table_plan, binding):
    if any(plan.get("ready") is not True or plan.get("accepted_new_runs") != 18 for plan in [main_plan, si_plan, table_plan]):
        raise ValueError("In-memory integration requires three complete eighteen-run plans")
    for plan in [main_plan, si_plan]:
        if digest(Path(plan["publication_values_path"])) != plan["publication_values_sha256"]:
            raise ValueError("Publication values changed after preparation")
    if main_plan["publication_values_sha256"] != si_plan["publication_values_sha256"]:
        raise ValueError("The main and SI plans use different publication values")
    files = set(main_plan["documents"]) | set(si_plan["documents"])
    documents, audit = {}, {}
    for name in sorted(files):
        document = Document(packet / name)
        citation_before = citations(document)
        marked = "Marked_Changes" in name
        before = normalized_text(document)
        if name in main_plan["documents"]:
            apply_main(document, main_plan, name, marked=marked)
        if name in si_plan["documents"]:
            apply_si(document, si_plan, name, marked=marked)
        if name.startswith("Supplementary_Information"):
            apply_updates(document, table_plan, binding, name, marked=marked)
        if citations(document) != citation_before:
            raise ValueError(f"Superscript citation sequence changed: {name}")
        audit[name] = {"input_docx_sha256": digest(packet / name), "citation_sequence_preserved": True,
                       "superscript_citation_run_count": len(citation_before), "accepted_text_changed": before != normalized_text(document)}
        documents[name] = document
    pairs = []
    for clean, marked in PAIRS:
        paragraph_match = normalized_text(documents[clean]) == normalized_text(documents[marked])
        table_match = accepted_tables(documents[clean]) == accepted_tables(documents[marked])
        if not paragraph_match or not table_match:
            raise ValueError(f"Accepted clean/marked documents differ after integration: {clean}")
        pairs.append({"clean": clean, "marked": marked, "accepted_paragraphs_match": paragraph_match,
                      "accepted_tables_match": table_match})
    return documents, {"schema_version": 1, "passed": True, "docx_saved": False,
                       "publication_values_sha256": main_plan["publication_values_sha256"], "documents": audit, "pairs": pairs}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--packet", type=Path, default=PACKET)
    parser.add_argument("--main-plan", type=Path, default=HERE / "main_result_replacements.json")
    parser.add_argument("--si-plan", type=Path, default=HERE / "si_result_replacements.json")
    parser.add_argument("--table-plan", type=Path, default=HERE / "si_result_table_updates.json")
    parser.add_argument("--binding", type=Path, default=HERE / "si_result_table_bindings.json")
    parser.add_argument("--audit", type=Path, default=HERE / "result_integration_check.json")
    args = parser.parse_args()
    documents, audit = integrate(args.packet, read_json(args.main_plan), read_json(args.si_plan), read_json(args.table_plan), read_json(args.binding))
    write_json(args.audit, audit)
    print(json.dumps({"passed": True, "documents": len(documents), "clean_marked_pairs": len(audit["pairs"]), "docx_saved": False}))


if __name__ == "__main__":
    main()
