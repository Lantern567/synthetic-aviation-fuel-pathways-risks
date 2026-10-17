from pathlib import Path
import argparse
import hashlib
import json
import re

import numpy as np
import pandas as pd
from docx import Document

from word_edits import accepted, replace_phrase
from update_si_results import PACKET, WORK, digest, read_json, write_json


HERE = Path(__file__).resolve().parent
SI_TARGETS = [27, 67, 85, 86, 95, 132, 135, 136, 140, 145, 146, 148, 152, 154, 158, 159, 161, 163, 167, 169,
              171, 175, 178, 179, 181, 183, 189, 192, 195, 196, 211, 215, 217, 218, 221, 222, 226,
              230, 234, 235, 237, 239, 240, 250, 255, 256, 257, 258, 259]
ARCHIVE_TARGETS = [4, 32, 34, 36, 38, 40, 42, 44]
FILES = {"Supplementary_Information.docx": SI_TARGETS, "Supplementary_Information_Marked_Changes.docx": SI_TARGETS,
         "Supplementary_Figure_Archive.docx": ARCHIVE_TARGETS, "Supplementary_Figure_Archive_Marked_Changes.docx": ARCHIVE_TARGETS}


def fmt(value, digits=2):
    if not np.isfinite(float(value)):
        raise ValueError("Publication values must be finite")
    return f"{float(value):.{digits}f}"


def inspect(packet):
    reference = {"si": Document(packet / "Supplementary_Information.docx"),
                 "archive": Document(packet / "Supplementary_Figure_Archive.docx")}
    documents = {}
    for name, targets in FILES.items():
        path = packet / name
        role = "archive" if "Archive" in name else "si"
        document = Document(path)
        lookup = {}
        for index, paragraph in enumerate(document.paragraphs):
            text = accepted(paragraph)
            lookup.setdefault(re.sub(r"\s+", "", text), []).append((index, text))
        records = {}
        for index in targets:
            expected = accepted(reference[role].paragraphs[index])
            matches = lookup.get(re.sub(r"\s+", "", expected), [])
            if len(matches) != 1:
                raise ValueError(f"Current SI result paragraph needs a unique match: {name}/P{index}")
            current_index, text = matches[0]
            records[str(index)] = {"paragraph_index": current_index, "before": text}
        documents[name] = {"source_path": str(path), "source_sha256": digest(path), "role": role, "paragraphs": records}
    return {"schema_version": 1, "docx_saved": False, "documents": documents,
            "root_owned_archive_statement": {"clean_si_paragraph_index": 244, "required_new_baselines": 12,
                                               "preserved_comparator": "GTL", "accepted_new_runs": 18,
                                               "fixed_asset_comparisons": 6, "week_omission_calculations": 52,
                                               "tag_and_commit": "Root supplies the final verified archive identifiers."}}


def load_values(path):
    data = read_json(path)
    if data.get("ready") is not True or data.get("baseline_count") != 13 or data.get("new_run_count") != 18:
        raise ValueError("SI results require the complete thirteen-baseline, eighteen-run publication values")
    for name, expected in data["source_sha256"].items():
        if digest(Path(name)) != expected:
            raise ValueError(f"Publication value source changed: {name}")
    return data


def bounds_record(path, publication):
    data = read_json(path)
    cases = {row["label"]: row for row in data["cases"]}
    if set(cases) != {"baseline", "summer_fixed"} or any(row["scenario"] != "CTL" for row in cases.values()):
        raise ValueError("CTL necessary-bound evidence must identify both actual CTL input sequences")
    for label, row in cases.items():
        if row["proves_original_90_percent_requirement_infeasible"] is not (label == "baseline"):
            raise ValueError("The CTL bound does not support the specified baseline/summer implication")
        for name, expected in row["input_sha256"].items():
            if digest(Path(name)) != expected:
                raise ValueError(f"CTL bound input changed: {name}")
    baseline = float(cases["baseline"]["necessary_resource_supply_fraction_upper_bound"])
    summer = float(cases["summer_fixed"]["necessary_resource_supply_fraction_upper_bound"])
    if not np.isclose(baseline, publication["ctl_baseline_necessary_resource_upper_bound"], atol=1e-6, rtol=0):
        raise ValueError("The CTL baseline bound differs from the accepted diagnostic scope")
    if not baseline < .9 < summer:
        raise ValueError("CTL baseline and summer necessary bounds must retain their distinct implications")
    return baseline, summer


def text_plan(values, costs, ranges_path, bound_path):
    b = values["baselines"]
    c, g, h, m = b["CCU-GH-FT"], b["GTL"], b["GTL-BH"], b["CCU-GH-MTJ"]
    ct, gt, ht = c["temporal_metrics"], g["temporal_metrics"], h["temporal_metrics"]
    cn, gn, hn = c["network_metrics"], g["network_metrics"], h["network_metrics"]
    sx = c["carbon_sensitivity"]
    pair = values["fixed_asset_comparisons"]
    a, z = pair["CCU-GH-FT"]["baseline"], pair["CCU-GH-FT"]["summer_fixed"]
    ga, gz = pair["GTL-GH"]["baseline"], pair["GTL-GH"]["summer_fixed"]
    baseline_bound, summer_bound = bounds_record(bound_path, values)
    grid, dac = pd.read_csv(costs / "S22_fixed_design_grid.csv"), pd.read_csv(costs / "S23_fixed_design_dac.csv")
    sensitivity = read_json(costs / "sensitivity_verification.json")
    share = float(sensitivity["S22"]["nonrenewable_output_share"])
    if not 0 <= share <= 1:
        raise ValueError("S22 nonrenewable-site output share is invalid")
    if len(grid) != 25 or set(dac.scenario) != {"DAC-GH-FT", "DAC-GH-MTJ", "DAC-BH-FT", "DAC-BH-MTJ"}:
        raise ValueError("Fixed-design sensitivity source tables are incomplete")
    dac_ranges = {}
    for scenario, frame in dac.groupby("scenario"):
        frame = frame.set_index("dac_cost_cny_t_co2")
        dac_ranges[scenario] = (float(frame.loc[300, "cost_cny_kg"]), float(frame.loc[4500, "cost_cny_kg"]))
    grid_range = (float(grid.cost_cny_kg.min()), float(grid.cost_cny_kg.max()))
    matching = pd.DataFrame(values["capacity_matching"])
    matching = matching[(matching.source == "GTL") & (matching.target == "CCU-GH-FT")].set_index("threshold_km")
    match25 = matching.loc[25]
    constant = np.allclose(matching.matched_capacity_kg_h, match25.matched_capacity_kg_h, atol=.01, rtol=0)
    ranges = read_json(ranges_path)
    outside = [int(row["week"]) for row in ranges if not row["inside"]]
    inside = sum(bool(row["inside"]) for row in ranges)
    if len(ranges) != 12 or sum(bool(row["selected"]) for row in ranges) != 4:
        raise ValueError("Joint-week input diagnostic is incomplete")
    scope = "twelve newly computed baseline configurations and the preserved GTL baseline"
    s, arc = {}, {}
    resource_path = WORK / "figures_resource/resource_data/selected_candidate_resource_metrics.csv"
    resource_metadata_path = WORK / "figures_resource/resource_data/resource_metrics.json"
    resource = pd.read_csv(resource_path)
    resource_metadata = read_json(resource_metadata_path)
    if resource_metadata["capacity_definition"] != "AC nominal power":
        raise ValueError("Resource-quality capacity definition differs from its actual AC denominator")
    raw_solar = float(resource.loc[resource.type.eq("solar_plant"), "capacity_raw_mw"].sum())
    ac_solar = float(resource.loc[resource.type.eq("solar_plant"), "capacity_mw"].sum())
    if not np.isclose(ac_solar, resource_metadata["solar"]["capacity_ac_mw"], atol=1e-8):
        raise ValueError("Resource-quality solar denominator differs from its metadata")
    s[27] = f"Raw solar ratings sum to {fmt(raw_solar, 1)} MW on a mixed DC/AC basis; the resource-quality metrics in Figs. S6–S7 use reconstructed AC ratings as their capacity denominator."
    s[67] = "Tables S6–S7 define 13 carbon–hydrogen–conversion cases. Twelve baseline configurations were recomputed, and GTL retains its archived baseline. Six electricity-dependent configurations also have fixed-production-asset alternative-week solutions, giving eighteen new solutions. CTL† denotes a partial-supply diagnostic with the original 90% output floor removed; its baseline resource bound establishes infeasibility with that floor. Cost and carbon intensity describe each configuration at its achieved demand-fulfillment level. Table S20 explains additional omitted combinations."
    s[85] = f"Capacity bounds, production requirements and scenario activation rules determine feasible scale. The CTL implementation includes a 90% minimum-production requirement; its baseline necessary resource upper bound is {fmt(baseline_bound * 100, 4)}%. The partial-supply diagnostic removes only that requirement, and both week sets use this diagnostic scope. The COPT calculations target a 1% gap for the full objective, including shortage penalties. Unit production cost is reported separately."
    s[86] = "Table S11 Configuration bounds and solver settings in the source archive. Table S27 reports COPT model dimensions, certified objective gaps and final-dispatch runtimes for the eighteen new solutions."
    s[95] = "Table S13 Unified constraint and accounting blocks used to express the 13 scenario-specific MILP instances. XE denotes ex-post accounting and is not a feasibility constraint; the remaining blocks define the feasible set in Eq. (S1). Section 5.1 and Tables S27–S29 identify the twelve newly computed baseline configurations, six alternative-week solutions and preserved GTL comparator."
    s[132] = "Twelve baseline configurations and six fixed-production-asset alternative-week solutions use COPT 8.0.7. GTL retains its archived Gurobi 12.0.3 baseline. A valid routing projection provides a global lower bound, candidate facility decisions and weekly route contracts are fixed, and the complete model is solved for period dispatch. Each accepted solution is feasible within its stated model scope and has a total-objective gap no greater than 1%. Table S27 reports the gaps, full-model sizes and final-dispatch runtimes; Tables S28–S29 report input, algebraic and independent physical checks."
    s[135] = "Table S16 describes the fixed-design parameter changes used for the supplementary cost comparisons. Figures S20–S24 reprice cost coefficients on the identified baseline designs, dispatch and output. Figures S20–S21 use the recomputed GTL-BH configuration and, where applicable, the preserved GTL baseline. Tables S21–S29 cover carbon accounting, utilization, site matching, week omission, joint input coverage and numerical validation. Alternative-week operating reoptimization is reported separately in Section 7.10."
    dr = "; ".join(f"{scenario}: {fmt(low)}–{fmt(high)} CNY/kg" for scenario, (low, high) in dac_ranges.items())
    s[136] = f"For CCU-GH-MTJ, varying grid price from 0.25 to 0.80 CNY/kWh and electrolyzer capital cost from 200,000 to 1,000,000 CNY/(kg H₂/h) gives {fmt(grid_range[0])}–{fmt(grid_range[1])} CNY/kg at fixed production (Fig. S22). Grid price affects synthesis at nonrenewable locations, which account for {fmt(share * 100)}% of fuel output in this solution. Changing DAC capture operating cost from 300 to 4,500 CNY/t gives the following fixed-design ranges (Fig. S23): {dr}. Figure S24 varies individual non-overlapping cost coefficients by ±20% and displays each pathway’s largest cost response. These calculations hold facility siting, dispatch, production and service fractions fixed."
    s[140] = "Tables S27–S29 document numerical validation of eighteen new COPT solutions: twelve baselines and six alternative-week solutions, including the two labelled CTL partial-supply diagnostics. Checks cover objective gaps, every exported constraint and variable, independent physical balances, input provenance and six baseline-to-alternative fixed-production-asset comparisons."
    s[145] = f"Over the baseline weeks, GTL-BH and GTL have worst-week retention of {fmt(ht['WorstWeekRetention'])}% and {fmt(gt['WorstWeekRetention'])}%, respectively. GTL meets each week’s airport demand; its lower output in calendar week 5 follows demand. CCU-GH-FT retains {fmt(ct['WorstWeekRetention'])}%. These values describe each selected schedule at its achieved output. Low-load durations describe within-week dispatch."
    s[146] = "Tables S23 and S25 report utilization metrics and static week-omission ranges for the twelve new baseline configurations and preserved GTL comparator. Section 7.10 evaluates the six electricity-dependent configurations with their baseline production assets fixed under the summer-inclusive input sequence. The six by-product-hydrogen cases and GTL contribute baseline schedules to the 13-pathway descriptive comparison."
    s[148] = f"GTL-BH and GTL cost {fmt(h['cost_cny_kg'])} and {fmt(g['cost_cny_kg'])} CNY/kg at demand coverage of {fmt(h['demand_fulfillment_ratio'] * 100)}% and {fmt(g['demand_fulfillment_ratio'] * 100)}%, respectively. CCU-GH-FT costs {fmt(c['cost_cny_kg'])} CNY/kg at {fmt(c['demand_fulfillment_ratio'] * 100)}% coverage. Tables S21–S22 report the costs, service levels and accounting variants for these configurations."
    s[152] = f"Figure S25 compares {scope}, with CTL identified as a partial-supply diagnostic. Cost and supply-chain carbon intensity are expressed per unit of actual output and accompanied by demand coverage. Tables S21–S22 provide carbon-accounting variants and comparison scope."
    s[154] = f"Figure S25. Extended cost and supply-chain carbon comparison across 13 pathway configurations. (a) Points retain pathway-family colors and FT/MTJ markers; the dashed line connects the nondominated baseline points. (b) Bars show costs excluding shortage penalties; end labels give supply-chain carbon intensity directly in g CO₂e/MJ. Aircraft combustion is excluded. The displayed envelope uses each configuration’s actual service fraction; Table S21 reports those fractions and the accounting sensitivity. The comparison includes {scope}; CTL is a partial-supply diagnostic."
    s[158] = f"Figure S27. Network and siting profiles within the retained pathway comparison groups. Active network identifiers refer to positive mass-flow endpoints; siting counts include selected facilities. The comparison includes {scope}, with CTL labelled as a partial-supply diagnostic. Group membership is retained for comparison."
    s[159] = "Table S18 Network descriptors calculated from the identified baseline transport records. GTL* is the preserved archived comparator; CTL† is the partial-supply diagnostic. Positive-flow identifiers, H₂ supply modules, synthesis modules and physical synthesis coordinates have distinct counting rules. Concentration metrics use mass flows; natural-gas volume is reported separately."
    service = f"CCU-GH-FT serves {fmt(c['demand_fulfillment_ratio'] * 100)}%, GTL {fmt(g['demand_fulfillment_ratio'] * 100)}% and GTL-BH {fmt(h['demand_fulfillment_ratio'] * 100)}% of baseline demand"
    s[161] = f"The cost-decomposition panel compares CCU-GH-FT, GTL and GTL-BH by cost category. Costs are normalized by actual fuel output and exclude shortage penalties. {service}. The decomposition describes their baseline production systems and achieved output."
    s[163] = f"Figure S28. Cost composition of three selected pathway configurations. Stacked bars partition reported discounted costs excluding shortage penalties into nonoverlapping categories. Cards give cost per kilogram of actual output and service fraction; {service}. Aggregate cost totals are excluded from component shares."
    s[167] = f"Figure S29. Positive production and supply-chain emissions of three selected pathway configurations. Stacked bars show positive emissions over four representative weeks; cards report the net ledger and capture attribution. {service}. Hydrogen and CO₂ transport subtotals are counted once. Aircraft combustion is excluded; Table S21 provides the accounting variants."
    s[169] = f"This section describes utilization over the baseline four-week sequence for {scope}, including the CTL partial-supply diagnostic. Utilization and worst-week retention characterize each schedule at its achieved output; demand coverage is reported separately."
    s[171] = "Facility utilization is computed for each three-hour production period as actual output divided by installed hourly capacity and period length. Hydrogen utilization uses selected electrolyzer capacity in the electrolysis configurations and selected conditioning capacity in the by-product-hydrogen configurations. The GTL comparator has no external hydrogen-production module and no H₂ utilization distribution."
    s[175] = f"Figure S30. Hydrogen-production and SynAF-synthesis utilization by feedstock category over 224 three-hour periods. Hydrogen denominators use selected electrolyzer or by-product-hydrogen conditioning capacity. Time-series curves show 24-h centred moving averages calculated separately within each representative week; distributions use unsmoothed three-hour values. The comparison includes {scope}, with CTL labelled as a partial-supply diagnostic."
    s[178] = f"Figure S31. Temporal comparisons for all 13 pathways: (a) worst-week retention versus total capacity premium, (b) worst-week-retention ranking, (c) mean-utilization and additional worst-week capacity premiums, and (d) longest low-load episode versus mean SynAF utilization. Episodes are counted within individual weeks. Records comprise {scope}, with CTL labelled as a partial-supply diagnostic."
    s[179] = "6.5.2 Utilization-based capacity-premium framework"
    s[181] = "The mean-utilization capacity premium expresses the extra nameplate synthesis capacity associated with producing the reported output at mean utilization:"
    s[183] = "The additional worst-week capacity premium expresses the further capacity associated with producing that output at the lowest weekly mean utilization:"
    s[189] = f"GTL-BH has {fmt(ht['SAFMean'])}% mean utilization and {fmt(ht['WorstWeekRetention'])}% worst-week retention. GTL has {fmt(gt['SAFMean'])}% mean utilization and {fmt(gt['WorstWeekRetention'])}% retention. CCU-GH-FT has {fmt(ct['SAFMean'])}% mean utilization and {fmt(ct['WorstWeekRetention'])}% retention, corresponding to mean-utilization and additional worst-week capacity premiums of {fmt(ct['ChronicPenaltyPct'])}% and {fmt(ct['RobustnessPenaltyPct'])}%. Its baseline demand coverage is {fmt(c['demand_fulfillment_ratio'] * 100)}%. Tables S19 and S23 report these properties, and Tables S25 and S32 describe sensitivity to week selection."
    s[192] = "Figure S33. Utilization distributions for GTL-BH, GTL and CCU-GH-FT, with H₂ supply at left and SynAF synthesis at right. Means and medians are annotated. GTL-BH uses selected hydrogen-conditioning capacity as its denominator; CCU-GH-FT uses installed electrolyzer capacity. Demand coverage is reported separately."
    s[195] = "6.5.4 Operating-metric fingerprint"
    s[196] = f"The twelve-metric fingerprint describes mean fuel utilization, within-week low-load frequency and duration, worst-week retention, interweek variability, renewable variability and capacity-premium components. Low-load events reset between nonconsecutive representative weeks; the more-than-12-hour count uses a strict threshold. Each value is standardized across {scope}, including the labelled CTL diagnostic. The heatmap describes the selected schedules at their achieved output."
    s[211] = "Table S21 Delivered-output costs, demand coverage and fixed-configuration carbon-accounting sensitivities. Carbon intensity is in g CO₂e/MJ. GTL* denotes the preserved archived comparator; CTL† denotes the partial-supply diagnostic."
    s[215] = f"A common combustion addition cancels from pairwise intensity differences and changes comparison with a combustion-inclusive benchmark. CCU-GH-FT has baseline intensity of {fmt(sx['baseline_g_mj'], 3)} g CO₂e/MJ, increasing to {fmt(sx['with_combustion_g_mj'], 3)} with combustion and {fmt(sx['with_combustion_without_capture_attribution_g_mj'], 3)} when capture attribution is also removed. The corresponding accounts hold the same configuration and output fixed."
    s[217] = "Aggregate utilization is total fuel output divided by installed synthesis capacity and modelled duration. With mean utilization Ū and the smallest weekly mean Umin expressed as fractions, the mean-utilization capacity premium is 100(1/Ū − 1), the additional worst-week capacity premium is 100(1/Umin − 1/Ū), and retention is 100Umin/Ū. These ratios quantify capacity requirements for the reported schedules. Demand coverage is reported separately."
    s[218] = "Table S23 Baseline utilization and implied capacity premiums. All numeric entries are percentages. GTL* denotes the preserved archived comparator; CTL† denotes the partial-supply diagnostic."
    s[221] = f"Synthesis modules with identical latitude and longitude are aggregated before matching. GTL, GTL-BH and CCU-GH-FT occupy {int(gn['saf_physical_sites'])}, {int(hn['saf_physical_sites'])} and {int(cn['saf_physical_sites'])} synthesis coordinates, respectively. Candidate pairs use great-circle distance. A maximum-weight one-to-one assignment credits the smaller installed capacity in each pair, with each coordinate used once. Dividing matched capacity by each pathway total gives the directional shares. Airport demand nodes are included only where actual synthesis is installed."
    threshold = "These shares are unchanged over the tested 1–100 km thresholds." if constant else "Table S24 reports the corresponding shares at each tested threshold from 1 to 100 km."
    s[222] = f"At 25 km, GTL and CCU-GH-FT have {match25.matched_capacity_kg_h:,.2f} kg/h of matched synthesis capacity: {fmt(match25.source_share_percent)}% of GTL capacity and {fmt(match25.target_share_percent)}% of CCU-GH-FT capacity. {threshold} These values measure geographic correspondence between selected configurations. Engineering reuse conditions are considered in Section 7.9."
    s[226] = "Each baseline week is omitted once while the other three weeks retain their installed capacity and dispatch. Thirteen baselines × four omissions give 52 descriptive calculations: 48 use the twelve new baseline schedules, including the labelled CTL diagnostic, and four use the preserved GTL schedule. Table S25 reports ranges with equal weighting of the retained weeks. Each calculation uses the remaining stored schedules without reoptimization."
    s[230] = f"For the 109 retained resource projects, weekly airport demand, solar energy and wind energy are normalized by their respective twelve-week means. The baseline four weeks define marginal minimum–maximum ranges. {inside} of the twelve candidate weeks lie inside all three ranges; weeks {', '.join(map(str, outside))} fall outside. This rectangular-range diagnostic identifies omitted demand–resource states in the candidate pool."
    s[234] = "Table S27 reports eighteen COPT 8.0.7 solutions: twelve baselines and six fixed-production-asset alternative-week solutions. Each complete-model dispatch is compared with a valid global lower bound from a routing projection. The relative gap applies to the complete objective, including shortage penalties. Candidate integers and H₂ and CO₂ weekly route contracts are re-evaluated under all period constraints of the stated model scope. Dimensions are before presolve. Times cover final complete-model dispatch only; accompanying records separately document design searches, projection solves and export."
    s[235] = "Table S27 COPT 8.0.7 certificates for eighteen new solutions. A and B denote baseline and summer-inclusive weeks; the six by-product-hydrogen configurations have A solutions only. CTL† removes the original 90% output floor in both A and B. Seconds cover final dispatch only."
    s[237] = "Table S28 Numerical and independent physical checks of eighteen accepted solutions. Constraint and variable residuals use exported model units; electricity deficits are in MWh and hydrogen balances in kg. Shared-renewable-budget electricity deficits apply to electrolysis configurations; grid-only by-product configurations show an em dash in that column."
    s[239] = "All exported linear constraints, variable bounds and integer decisions are checked. Independent calculations reconstruct source supplies, electricity use, hydrogen and fuel balances, capacity limits, initial inventories, weekly demand and route volumes. DAC additionally has node-period, source-week and system-wide CO₂ checks. GTL-GH A/B and GTL-BH A also pass 28 daily natural-gas allocations, source-technology compatibility, truck and LNG-terminal capacities and transport/upstream ledger checks. Explicit transport and airport-integrated gas procurement are accounted for separately. Cost and emissions checks use actual four-week transport and 672-hour inventory accounting."
    s[240] = "Table S29 records independent input, physical-accounting and fixed-asset checks. All eighteen new solutions meet the stated tolerances within their identified model scopes. The certified gap applies to the full objective, including shortage penalties; reported unit fuel costs exclude those penalties."
    selected_match = matching.loc[[25, 50, 100]]
    if np.allclose(selected_match.source_share_percent, selected_match.source_share_percent.iloc[0]) and np.allclose(selected_match.target_share_percent, selected_match.target_share_percent.iloc[0]):
        matchtext = f"At 25, 50 and 100 km, matched synthesis capacity represents {fmt(match25.source_share_percent)}% of GTL capacity and {fmt(match25.target_share_percent)}% of CCU-GH-FT capacity."
    else:
        matchtext = " ".join(f"At {int(radius)} km, matched capacity represents {fmt(row.source_share_percent)}% of GTL and {fmt(row.target_share_percent)}% of CCU-GH-FT capacity." for radius, row in selected_match.iterrows())
    s[250] = f"GTL and CCU-GH-FT both use Fischer–Tropsch conversion. {matchtext} These shares provide geographic screening within the same broad conversion family. Reactor, catalyst and feed-conditioning compatibility require equipment-specific information. GTL-BH uses methanol-to-jet, so comparison with an FT route also requires assessment of conversion-process changes. Downstream storage and delivery are assessed separately under the asset-specific criteria in Table S31."
    s[255] = "A: weeks 1, 5, 14 and 44. B: weeks 1, 14, 27 and 44. Utilization and premiums follow Section 7.3. Demand coverage is delivered fuel divided by selected demand. CTL† removes the original 90% minimum-production inequality in both sets; B fixes the production assets from diagnostic A."
    demand_a, demand_z = a["demand_kg"], z["demand_kg"]
    supply = "; ".join(f"{name}: {fmt(records['baseline']['demand_fulfillment_ratio'] * 100)}% / {fmt(records['summer_fixed']['demand_fulfillment_ratio'] * 100)}%" for name, records in pair.items())
    s[256] = f"Selected demand increases by {fmt((demand_z / demand_a - 1) * 100)}%, from {fmt(demand_a / 1e6, 3)} to {fmt(demand_z / 1e6, 3)} million kg. Demand coverage in A/B is {supply}. Costs exclude shortage penalties and describe actual production."
    s[257] = f"For CCU-GH-FT, mean synthesis utilization is {fmt(a['mean_utilization'] * 100)}% in A and {fmt(z['mean_utilization'] * 100)}% in B; worst-week retention is {fmt(a['worst_week_retention'] * 100)}% and {fmt(z['worst_week_retention'] * 100)}%. The mean-utilization capacity premium is {fmt(a['mean_utilization_capacity_premium_pct'])}% and {fmt(z['mean_utilization_capacity_premium_pct'])}%, and the additional worst-week premium is {fmt(a['additional_worst_week_capacity_premium_pct'])}% and {fmt(z['additional_worst_week_capacity_premium_pct'])}%. GTL-GH worst-week retention is {fmt(ga['worst_week_retention'] * 100)}% in A and {fmt(gz['worst_week_retention'] * 100)}% in B. These calculations quantify selected-week operating sensitivity for the six fixed production designs. The six by-product configurations and preserved GTL comparator contribute baseline schedules only."
    s[258] = f"CTL’s necessary renewable-energy-only production upper bounds are {fmt(baseline_bound * 100, 4)}% for A and {fmt(summer_bound * 100, 4)}% for B, before synthesis electricity, logistics and installed-capacity restrictions. The A bound is below 90% and establishes infeasibility of the baseline formulation with that floor. The B bound exceeds 90% and does not establish infeasibility for its week set. Both CTL† results use the diagnostic formulation with the floor removed, with B fixing the production assets selected in diagnostic A."
    s[259] = "All twelve A/B solutions pass algebraic and independent physical checks within the stated model scopes, and all six pairs pass the fixed-production-asset comparison. Table S27 also records the six new by-product-hydrogen baseline solutions. It reports model dimensions, final-dispatch times and certified objective gaps. Both week sets use the same shortage coefficient and process cost assumptions."
    arc[4] = f"This companion presents input summaries, fixed-design cost sensitivities, conditional learning projections and pathway maps for “Infrastructure incompatibility across synthetic aviation fuel pathways risks stranding decarbonization investments”. Figure numbers are shared with the core Supplementary Information. The baseline comparison includes {scope}, with CTL labelled as a partial-supply diagnostic. Six electricity-dependent configurations also have fixed-production-asset alternative-week solutions. The core supplement reports model assumptions, demand coverage and numerical validation. Conditional cost projections hold baseline designs, dispatch and production fixed."
    arc[32] = "Figure S20. Fixed-design natural-gas-price sensitivity for the preserved GTL baseline and recomputed GTL-BH configuration. The natural-gas procurement coefficient is repriced using each saved design, actual gas purchases, dispatch and fuel output. The shaded band marks the illustrative 6–8 CNY/kg market-price window."
    arc[34] = "Figure S21. Fixed-design by-product-hydrogen conditioning-capital sensitivity for the recomputed GTL-BH configuration. The horizontal axis varies the steel-source PSA-plus-compression CAPEX in CNY/(kg H₂/h), with the refinery coefficient maintained at 0.8 times that value. Repricing uses selected conditioning capacities and the saved capital-accounting coefficients, with dispatch and actual fuel output fixed. The shaded band marks 6–8 CNY/kg; dashed guides identify the configured steel and refinery coefficients, 280,000 and 224,000 CNY/(kg H₂/h)."
    arc[36] = f"Figure S22. Fixed-design cost sensitivity of CCU-GH-MTJ to grid electricity price and electrolyzer capital cost. Twenty-five entries reprice grid electricity at 0.25–0.80 CNY/kWh and raw electrolyzer CAPEX at 200,000–1,000,000 CNY/(kg H₂/h), holding facilities, dispatch, output and {fmt(m['demand_fulfillment_ratio'] * 100)}% service fixed. The baseline pair is 0.60 CNY/kWh and 360,000 CNY/(kg H₂/h). Unit costs range from {fmt(grid_range[0])} to {fmt(grid_range[1])} CNY/kg. Grid-price changes apply to synthesis at nonrenewable sites; renewable-hydrogen electricity retains its baseline coefficient."
    arc[38] = f"Figure S23. Fixed-design DAC-linked cost sensitivity to capture operating cost from 300 to 4,500 CNY/t CO₂. The four curves use the recomputed DAC-GH-FT, DAC-GH-MTJ, DAC-BH-FT and DAC-BH-MTJ designs and actual captured quantities. Costs, facilities, dispatch and output otherwise remain fixed. Endpoint costs are {dr}. The 100 CNY/t vertical guide is an illustrative value outside the sweep. Service fractions are retained at their respective baselines."
    arc[40] = f"Figure S24. Largest fixed-design unit-cost response to one-component-at-a-time ±20% perturbations. Each positive, nonoverlapping cost component is multiplied by 0.8 and 1.2, holding the other coefficients, facilities, dispatch, output and service fixed. Bars give the largest response for each of {scope}; CTL is a partial-supply diagnostic. Diamonds identify the GTL and GTL-BH baselines, and the shaded band marks 6–8 CNY/kg."
    arc[42] = f"Figure S26. Joint economic, spatial and temporal comparisons for {scope}, with CTL labelled as a partial-supply diagnostic. Network statistics use mean weekly mass flows, with natural-gas volumes reported separately. Siting shares count selected SynAF modules. Panel D gives actual-output costs and supply-chain carbon intensities directly in CNY/kg and g CO₂e/MJ. Metrics describe the solved configurations and their actual service levels."
    arc[44] = f"Figure S36. Fixed-design cost projection and positive supply-chain carbon composition for CTL. Polar bars show 2026–2036 category costs under the common learning assumptions; the outer ring and carbon donut show positive baseline emissions before credits, in g CO₂e/MJ. The cost donut partitions baseline expenditure excluding shortage penalties. Design, dispatch, output and service remain fixed; aircraft combustion is excluded. CTL uses the partial-supply diagnostic at {fmt(b['CTL']['demand_fulfillment_ratio'] * 100)}% service. Table S21 reports service and carbon-accounting variants."
    tracked = {str(costs / name): digest(costs / name) for name in ["S22_fixed_design_grid.csv", "S23_fixed_design_dac.csv", "sensitivity_verification.json"]}
    tracked.update({str(ranges_path): digest(ranges_path), str(bound_path): digest(bound_path)})
    tracked.update({str(resource_path): digest(resource_path), str(resource_metadata_path): digest(resource_metadata_path)})
    return {"si": s, "archive": arc}, tracked


def prepare(locations, values, replacements, tracked, values_path):
    documents = {}
    for name, record in locations["documents"].items():
        edits = []
        for key, binding in record["paragraphs"].items():
            before, after = binding["before"], replacements[record["role"]][int(key)]
            if record["role"] == "si" and int(key) == 27:
                after = before if after in before else before + " " + after
            if record["role"] == "si" and int(key) in [171, 196]:
                citation = re.search(r"\d+(?:,\d+)*$", before)
                if citation:
                    before = before[:citation.start()]
            if before != after:
                edits.append({"paragraph_index": binding["paragraph_index"], "clean_reference_index": int(key), "before": before, "after": after,
                              "paragraph_before_sha256": hashlib.sha256(binding["before"].encode()).hexdigest(),
                              "paragraph_after": binding["before"].replace(before, after, 1)})
        documents[name] = {"source_path": record["source_path"], "source_sha256": record["source_sha256"], "edits": edits}
    return {"schema_version": 1, "ready": True, "docx_saved": False, "accepted_new_runs": 18,
            "publication_values_path": str(values_path), "publication_values_sha256": digest(values_path),
            "source_sha256": {**values["source_sha256"], **tracked}, "documents": documents,
            "root_owned_archive_statement": locations["root_owned_archive_statement"]}


def apply_replacements(document, plan, filename, marked=False):
    if plan.get("ready") is not True or plan.get("accepted_new_runs") != 18:
        raise ValueError("SI edits require the complete accepted result plan")
    if digest(Path(plan["publication_values_path"])) != plan["publication_values_sha256"]:
        raise ValueError("Prepared publication values changed")
    for path, expected in plan["source_sha256"].items():
        if digest(Path(path)) != expected:
            raise ValueError(f"Prepared SI source changed: {path}")
    for edit in plan["documents"][filename]["edits"]:
        paragraph = document.paragraphs[edit["paragraph_index"]]
        if hashlib.sha256(accepted(paragraph).encode()).hexdigest() != edit["paragraph_before_sha256"]:
            raise ValueError(f"Current SI paragraph changed: {filename}/P{edit['paragraph_index']}")
        replace_phrase(paragraph, edit["before"], edit["after"], marked=marked)
        if accepted(paragraph) != edit["paragraph_after"]:
            raise ValueError("Accepted SI result replacement differs")
    return document


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["inspect", "prepare"])
    parser.add_argument("--packet", type=Path, default=PACKET)
    parser.add_argument("--locations", type=Path, default=HERE / "si_result_locations.json")
    parser.add_argument("--out", type=Path, default=HERE / "si_result_replacements.json")
    parser.add_argument("--values", type=Path, default=HERE / "publication_values.json")
    parser.add_argument("--costs", type=Path, default=WORK / "figures_cost/adapted_data")
    parser.add_argument("--ranges", type=Path, default=WORK / "figures_resource/updated_joint_week_ranges.json")
    parser.add_argument("--ctl-bounds", type=Path, default=WORK / "validation/ctl_resource_bounds.json")
    args = parser.parse_args()
    locations = inspect(args.packet)
    if args.mode == "inspect":
        write_json(args.locations, locations)
        print(json.dumps({"mode": "inspect", "documents": len(locations["documents"]), "docx_saved": False}))
    else:
        values = load_values(args.values)
        replacements, tracked = text_plan(values, args.costs, args.ranges, args.ctl_bounds)
        plan = prepare(locations, values, replacements, tracked, args.values)
        write_json(args.out, plan)
        print(json.dumps({"mode": "prepare", "documents": len(plan["documents"]), "docx_saved": False}))


if __name__ == "__main__":
    main()
